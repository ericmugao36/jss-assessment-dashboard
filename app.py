import os
import json
import html
import math
import secrets
from datetime import date, datetime
from io import BytesIO
from tracemalloc import start

from dotenv import load_dotenv
from fastapi import FastAPI, Form, Request, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse
from starlette.middleware.sessions import SessionMiddleware

from sqlalchemy import (
    create_engine, String, Integer, Float, Date, DateTime,
    ForeignKey, Table, Column, Text, select, event
)
from sqlalchemy.orm import (
    DeclarativeBase, Mapped, mapped_column, relationship, sessionmaker
)
from pwdlib import PasswordHash
from openpyxl import load_workbook


# ============================================================
# 1. CONFIGURATION
# ============================================================

load_dotenv(override=True)

SCHOOL = os.getenv("SCHOOL_NAME", "Your School")
print("SCHOOL NAME LOADED:", SCHOOL)

DATABASE_URL = os.getenv(
    "DATABASE_URL",
    "sqlite:///./school.db",
)

if DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace(
        "postgres://",
        "postgresql+psycopg://",
        1,
    )
elif DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace(
        "postgresql://",
        "postgresql+psycopg://",
        1,
    )

SECRET = os.getenv("SESSION_SECRET", "")
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "false").lower() == "true"
APP_ENV = os.getenv("APP_ENV", "development").lower()

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 20000

SUBJECTS = [
    "English",
    "Kiswahili",
    "Mathematics",
    "Kenyan Sign Language",
    "Integrated Science",
    "Agriculture",
    "Social Studies",
    "CRE/IRE/HRE",
    "Creative Arts & Sports",
    "Pre-technical Studies",
    "Business Studies",
    "Life Skills Education",
]

GRADES = list(range(1, 13))

DEFAULT_SHEETS = {
    "learners_sheet": "Learners",
    "parents_sheet": "Parents",
    "links_sheet": "ParentLinks",
    "marks_sheet": "AssessmentMarks",
}

if APP_ENV == "production":
    if len(SECRET) < 32:
        raise RuntimeError(
            "Set a strong SESSION_SECRET in your .env file."
        )
    if not COOKIE_SECURE:
        raise RuntimeError(
            "Set COOKIE_SECURE=true when using HTTPS in production."
        )
elif len(SECRET) < 32:
    SECRET = secrets.token_urlsafe(48)
    print(
        "WARNING: Temporary session secret. "
        "Set SESSION_SECRET in your .env file."
    )


 # SQLite waits for busy writers and uses WAL for better
# reader/writer concurrency. It still allows only one writer
# at a time.
engine = create_engine(
    DATABASE_URL,
    connect_args={
        "check_same_thread": False,
        "timeout": 30,
    } if DATABASE_URL.startswith("sqlite") else {},
    pool_pre_ping=True,
)

if DATABASE_URL.startswith("sqlite"):

    @event.listens_for(engine, "connect")
    def set_sqlite_options(dbapi_connection, connection_record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA busy_timeout=30000")
        cursor.close()

    with engine.connect() as connection:
        connection.exec_driver_sql("PRAGMA journal_mode=WAL")
        connection.commit()


DB = sessionmaker(bind=engine, expire_on_commit=False)
passwords = PasswordHash.recommended()

app = FastAPI(title="EduTrack Kenya")

app.add_middleware(
    SessionMiddleware,
    secret_key=SECRET,
    session_cookie="jss_session",
    max_age=28800,
    same_site="lax",
    https_only=COOKIE_SECURE,
)


# ============================================================
# 2. DATABASE MODELS
# ============================================================

class Base(DeclarativeBase):
    pass


class School(Base):
    __tablename__ = "schools"

    id: Mapped[int] = mapped_column(primary_key=True)

    name: Mapped[str] = mapped_column(
        String(200),
        nullable=False,
        index=True,
    )

    # Maximum number of active teachers allowed.
    teacher_limit: Mapped[int] = mapped_column(
        Integer,
        default=100,
        nullable=False,
    )

    # Possible values: active, suspended.
    status: Mapped[str] = mapped_column(
        String(20),
        default="active",
        nullable=False,
    )

    created_at: Mapped[datetime] = mapped_column(
        DateTime,
        default=datetime.utcnow,
        nullable=False,
    )


class SchoolClass(Base):
    __tablename__ = "school_classes"

    id: Mapped[int] = mapped_column(primary_key=True)

    school_id: Mapped[int] = mapped_column(
        ForeignKey("schools.id"),
        nullable=False,
        index=True,
    )

    name: Mapped[str] = mapped_column(
        String(100),
        nullable=False,
    )

    grade: Mapped[int] = mapped_column(
        Integer,
        nullable=False,
    )


class TeacherClassAssignment(Base):
    __tablename__ = "teacher_class_assignments"

    id: Mapped[int] = mapped_column(primary_key=True)

    teacher_id: Mapped[int] = mapped_column(
        ForeignKey("users.id"),
        nullable=False,
        index=True,
    )

    class_id: Mapped[int] = mapped_column(
        ForeignKey("school_classes.id"),
        nullable=False,
        index=True,
    )




parent_links = Table(
    "parent_links",
    Base.metadata,
    Column("parent_id", ForeignKey("users.id"), primary_key=True),
    Column("student_id", ForeignKey("students.id"), primary_key=True),
)



class User(Base):
    __tablename__ = "users"

    id: Mapped[int] = mapped_column(primary_key=True)
    school_id: Mapped[int | None] = mapped_column(
        ForeignKey("schools.id"),
        nullable=True,
        index=True,
    )
    name: Mapped[str | None] = mapped_column(
        String(160), nullable=True
    )
    email: Mapped[str] = mapped_column(
        String(255), unique=True, index=True
    )
    password_hash: Mapped[str] = mapped_column(String(255))
    role: Mapped[str] = mapped_column(String(20))

    children: Mapped[list["Student"]] = relationship(
        secondary=parent_links,
        back_populates="parents",
    )


class Student(Base):
    __tablename__ = "students"

    id: Mapped[int] = mapped_column(primary_key=True)

    name: Mapped[str] = mapped_column(String(160))

    grade: Mapped[int] = mapped_column(Integer)

    admission: Mapped[str | None] = mapped_column(
        String(80),
        unique=True,
        index=True,
        nullable=True,
    )

    school_id: Mapped[int | None] = mapped_column(
        ForeignKey("schools.id"),
        nullable=True,
        index=True,
    )
    parents: Mapped[list[User]] = relationship(
        secondary=parent_links,
        back_populates="children",
    )

    results: Mapped[list["Result"]] = relationship(
        back_populates="student",
        cascade="all, delete-orphan",
    )


class Result(Base):
    __tablename__ = "results"

    id: Mapped[int] = mapped_column(primary_key=True)
    student_id: Mapped[int] = mapped_column(
        ForeignKey("students.id"), index=True
    )
    subject: Mapped[str] = mapped_column(String(120))
    assessment: Mapped[str] = mapped_column(String(120))
    result_date: Mapped[date] = mapped_column(Date)
    marks: Mapped[float] = mapped_column(Float)
    maximum: Mapped[float] = mapped_column(Float, default=100)

    student: Mapped[Student] = relationship(
        back_populates="results"
    )

class TeacherRemark(Base):
    __tablename__ = "teacher_remarks"

    id: Mapped[int] = mapped_column(primary_key=True)

    student_id: Mapped[int] = mapped_column(
        ForeignKey("students.id", ondelete="CASCADE"),
        index=True,
    )

    teacher_id: Mapped[int] = mapped_column(
        ForeignKey("users.id"),
        index=True,
    )

    category: Mapped[str] = mapped_column(
        String(80),
        default="Academic progress",
    )

    comment: Mapped[str] = mapped_column(Text)

    created_at: Mapped[datetime] = mapped_column(
        DateTime,
        default=datetime.utcnow,
    )

    student: Mapped["Student"] = relationship()
    teacher: Mapped["User"] = relationship()


class AppSetting(Base):
    __tablename__ = "app_settings"

    key: Mapped[str] = mapped_column(
        String(100), primary_key=True
    )
    value: Mapped[str] = mapped_column(String(255))


class ImportBatch(Base):
    __tablename__ = "import_batches"

    token: Mapped[str] = mapped_column(
        String(100), primary_key=True
    )
    teacher_id: Mapped[int] = mapped_column(ForeignKey("users.id"))
    payload: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(
        DateTime, default=datetime.utcnow
    )


Base.metadata.create_all(engine)

# ============================================================
#2b MULTI-SCHOOL MIGRATION: USER SCHOOL ASSOCIATION
# ============================================================

from sqlalchemy import inspect, text

with engine.begin() as connection:
    inspector = inspect(connection)

    # Add school_id only if it does not already exist.
    user_columns = {
        column["name"]
        for column in inspector.get_columns("users")
    }

    if "school_id" not in user_columns:
        connection.execute(
            text(
                "ALTER TABLE users "
                "ADD COLUMN school_id INTEGER "
                "REFERENCES schools(id)"
            )
        )

    # Create the initial configured school record if needed.
    school_id = connection.execute(
        text(
            "SELECT id FROM schools "
            "WHERE name = :name LIMIT 1"
        ),
        {"name": SCHOOL},
    ).scalar_one_or_none()

    if school_id is None:
        connection.execute(
            text(
                """
               INSERT INTO schools (
    name,
    teacher_limit,
    status,
    created_at
)
VALUES (
    :name,
    :teacher_limit,
    :status,
    :created_at
)
                """
            ),
            {
                "name": SCHOOL,
                "teacher_limit": 100,
                "status": "active",
                "created_at": datetime.utcnow(),
            },
        )

        school_id = connection.execute(
            text(
                "SELECT id FROM schools "
                "WHERE name = :name LIMIT 1"
            ),
            {"name": SCHOOL},
        ).scalar_one()

    # Associate existing users with the configured school.
    connection.execute(
        text(
            "UPDATE users SET school_id = :school_id "
            "WHERE school_id IS NULL"
        ),
        {"school_id": school_id},
    )

    # Add school_id to existing students table if missing.
    student_columns = {
        column["name"]
        for column in inspect(connection).get_columns("students")
    }

    if "school_id" not in student_columns:
        connection.execute(
            text(
                "ALTER TABLE students "
                "ADD COLUMN school_id INTEGER "
                "REFERENCES schools(id)"
            )
        )

    # Assign existing learners to the configured school.
    connection.execute(
        text(
            "UPDATE students SET school_id = :school_id "
            "WHERE school_id IS NULL"
        ),
        {"school_id": school_id},
    )


# ============================================================
# 3. HELPER FUNCTIONS
# ============================================================

def esc(value):
    return html.escape(
        str(value if value is not None else ""),
        quote=True,
    )


def go(path):
    return RedirectResponse(path, status_code=303)


def csrf(request: Request):
    token = request.session.get("csrf_token")

    if not token:
        token = secrets.token_urlsafe(32)
        request.session["csrf_token"] = token

    return token


def check_csrf(request: Request, supplied: str):
    expected = request.session.get("csrf_token", "")

    if not expected or not secrets.compare_digest(
        expected, supplied or ""
    ):
        raise HTTPException(
            status_code=403,
            detail="Your form expired. Refresh the page and try again.",
        )


def current_user(request: Request, db):
    user_id = request.session.get("user_id")

    if not user_id:
        return None

    return db.get(User, int(user_id))



def require_role(request: Request, db, role: str):
    user = current_user(request, db)

    if not user:
        raise HTTPException(
            status_code=303,
            headers={"Location": "/login"},
        )

    if user.role != role:
        raise HTTPException(
            status_code=403,
            detail="You do not have permission to access this page.",
        )

    # Teachers and parents must belong to an active school.
    if role in {"teacher", "parent"}:
        if not user.school_id:
            raise HTTPException(
                status_code=403,
                detail="Your account is not assigned to a school.",
            )

        school = db.get(School, user.school_id)

        if not school or school.status != "active":
            raise HTTPException(
                status_code=403,
                detail="Your school is inactive or unavailable.",
            )

    return user



def get_setting(db, key, default):
    row = db.get(AppSetting, key)
    return row.value if row else default


def set_setting(db, key, value):
    row = db.get(AppSetting, key)

    if row:
        row.value = value
    else:
        db.add(AppSetting(key=key, value=value))


def page(request, title, body, user=None, status_code=200):
    """Render every page with the shared EduTrack Kenya layout."""
    destination = "/teacher" if user and user.role == "teacher" else "/parent"
    sidebar = ""

    if user:
        teacher_links = ""
        if user.role == "teacher":
            teacher_links = f"""
                <a class="side-link" href="/import">📥 <span>Excel Import</span></a>
                <a class="side-link" href="/import/settings">⚙️ <span>Import Settings</span></a>
            """

        sidebar = f"""
        <aside class="sidebar">
            <div class="side-label">NAVIGATION</div>
            <a class="side-link" href="{destination}">▦ <span>Home</span></a>
{(
    '<a class="side-link" href="/teacher/add-learner">'
    '➕ <span>Add Learner</span></a>'
    '<a class="side-link" href="/teacher/learners">'
    '♙ <span>All Learners</span></a>'
    if user.role == "teacher" else ""
)}
            {teacher_links}
            <a class="side-link" href="/settings">⚙ <span>Settings</span></a>
            <div class="side-spacer"></div>
            <form method="post" action="/logout" class="logout-form">
                <input type="hidden" name="csrf_token" value="{esc(csrf(request))}">
                <button class="side-link logout-button" type="submit">↪ <span>Log out</span></button>
            </form>
        </aside>
        """

    document = f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)} - {esc(SCHOOL)}</title>


<style>
:root {{
  color-scheme: light;
  --page: #edf2f7; --surface: #ffffff; --surface-alt: #e4ebf4;
  --text: #172338; --muted: #617087; --border: #dbe3ef;
  --navy: #173b68; --navy-hover: #245a81; --link: #225da0;
  --sidebar: #e8eef7; --shadow: 0 4px 16px rgba(30,55,90,.06);
}}
body.dark-theme {{
  color-scheme: dark;
  --page: #080b12; --surface: #101722; --surface-alt: #1e293b;
  --text: #edf3fc; --muted: #b5c1d3; --border: #2b3545;
  --navy: #172957; --navy-hover: #294478; --link: #8bbcff;
  --sidebar: #1e293b; --shadow: 0 4px 16px rgba(0,0,0,.18);
}}

html.dark-start {{
  color-scheme: dark;
  background: #080b12;
}}

html.dark-start body {{
  background: #080b12;
}}


* {{ box-sizing:border-box; }}
body {{
    height: 100vh;
    overflow: hidden;
}}
body {{ margin:0; background:var(--page); color:var(--text); font:15px Arial,sans-serif; line-height:1.5; transition:background .2s,color .2s; }}

.topbar {{
    position: sticky;
    top: 0;
    z-index: 1000;
    min-height: 54px;
    margin: 0 10px;
    padding: 10px 16px;
    border-radius: 0 0 9px 9px;
    background: var(--navy);
    color: white;
    display: flex;
    justify-content: space-between;
    align-items: center;
    gap: 12px;
}}

.brand {{ color:white; text-decoration:none; font-size:17px; font-weight:bold; display:flex; gap:10px; align-items:center; }}
.brand-mark {{ font-size:22px; }}
.top-actions {{ display:flex; gap:14px; align-items:center; }}
.top-actions a {{ color:white; text-decoration:none; }}

.app-shell {{
    display: block;
    margin: 6px 10px 10px 232px;
}}
/* Dark blue styling for the Total Learners widget */
.total-learners-widget .flip-card-front {{
    background: #12345A;
    color: #FFFFFF;
}}

.total-learners-widget .flip-card-front h2,
.total-learners-widget .flip-card-front p,
.total-learners-widget .flip-card-front .stat-label,
.total-learners-widget .flip-card-front .stat-value {{
    color: #FFFFFF;
}}


.sidebar {{
    position: fixed;
    top: 70px;
    left: 10px;
    bottom: 10px;
    width: 210px;
    overflow-y: auto;
    background: var(--sidebar);
    border: 1px solid var(--border);
    border-radius: 12px;
    padding: 14px 10px;
    display: flex;
    flex-direction: column;
    gap: 5px;
    z-index: 900;
}}

.side-label {{ font-weight:bold; font-size:13px; margin:0 0 7px 7px; color:var(--text); }}
.side-link {{ display:flex; align-items:center; gap:10px; padding:9px 8px; border-radius:6px; text-decoration:none; color:var(--text); font-size:13px; border:0; background:transparent; width:100%; text-align:left; cursor:pointer; font-family:inherit; }}
.side-link:hover {{ background:rgba(95,135,190,.16); }}
.side-spacer {{ flex:1; min-height:18px; }}
.logout-form {{ margin:0; }}
main {{ min-width:0; padding:9px 0 20px; }}
main h1 {{ font-size:25px; margin:0 0 8px; }}
h2 {{ font-size:20px; }}
section {{ background:var(--surface); color:var(--text); padding:14px; margin:8px 0; border:1px solid var(--border); border-radius:10px; box-shadow:var(--shadow); overflow:hidden; }}

.dashboard-widgets {{
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    margin: 12px 0;
}}

.dashboard-widgets section {{
    margin: 0;
    min-width: 0;
}}

.total-learners-widget .stat-card {{
    min-height: 125px;
}}

.summary-label {{
    color: var(--muted);
    font-size: 13px;
    margin: 0 0 8px;
}}

.total-learners-widget .stat-card h2 {{
    font-size: 25px;
    overflow-wrap: anywhere;
    margin: 5px 0 10px;
}}

.chart-widget {{
    grid-column: 1 / -1;
}}

.chart-widget > p {{
    color: var(--muted);
    font-size: 13px;
}}

.subject-chart {{
    display: grid;
    grid-template-columns: 48px minmax(0, 1fr);
    gap: 10px;
    margin-top: 22px;
}}

.chart-y-axis {{
    box-sizing: border-box;
    height: 177px;
    min-height: 177px;
    padding: 0 0 27px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    align-items: flex-end;
    font-size: 12px;
    line-height: 1.2;
    color: var(--muted);
}}


.chart-plot {{
    box-sizing: border-box;
    width: 100%;
    min-width: 360px;
    height: 177px;
    display: flex;
    align-items: stretch;
    gap: 10px;
    background:
        repeating-linear-gradient(
            to top,
            transparent 0,
            transparent calc(20% - 1px),
            var(--border) calc(20% - 1px),
            var(--border) 20%
        );
    background-size: 100% 150px;
    background-position: left bottom;
    background-repeat: no-repeat;
    border-bottom: 1px solid var(--border);
}}


.subject-column {{
    flex: 1 1 0;
    min-width: 58px;
    height: 100%;
    position: relative;
}}


.subject-plot-column {{
    height: 150px;
    min-height: 150px;
    position: relative;
    box-sizing: border-box;
}}

.subject-bar {{
    position: absolute;
    bottom: 0;
    left: 25%;
    width: 50%;
    background: #3b82f6;
    border-radius: 5px 5px 0 0;
    min-height: 0;
    box-sizing: border-box;
    border: 1px solid #1d4ed8;
    box-sizing: border-box;
}}


.subject-value {{
    position: absolute;
    left: 0;
    width: 100%;
    text-align: center;
    font-size: 11px;
    font-weight: bold;
    color: var(--text);
    white-space: nowrap;
}}

.subject-name {{
    height: 27px;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
    text-align: center;
    font-size: 11px;
    color: var(--muted);
    padding-top: 5px;
}}

.chart-x-title,
.chart-y-title {{
    text-align: center;
    font-size: 12px;
    color: var(--muted);
}}

@media (max-width: 600px) {{
    .dashboard-widgets {{
        grid-template-columns: minmax(0, 1fr);
    }}

    .chart-widget {{
        grid-column: auto;
    }}
}}

form.grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:13px; align-items:end; }}
label {{ display:block; font-weight:bold; font-size:13px; }}
input,select,textarea {{ display:block; width:100%; padding:10px; margin-top:5px; border:1px solid var(--border); border-radius:7px; background:var(--surface); color:var(--text); font:inherit; }}
[hidden] {{ display: none !important;}}
button,.button,.dashboard-button {{ display:inline-block; padding:10px 13px; background:var(--navy); color:white; border:0; border-radius:7px; cursor:pointer; text-decoration:none; font:inherit; }}
button:hover,.button:hover,.dashboard-button:hover {{ background:var(--navy-hover); }}
.dashboard-button {{ margin:5px; font-weight:bold; }}
a {{ color:var(--link); }}
.topbar a {{ color:white; }}
table {{ border-collapse:collapse; width:100%; min-width:0; table-layout:fixed; }}
th, td {{ padding:8px; overflow-wrap:anywhere; }}
th,td {{ padding:6px 8px; text-align:left; border-bottom:1px solid var(--border); }}
th {{ background:var(--surface-alt); }}
.table {{ overflow-x:auto; margin-top:15px; }}
.note {{ background:var(--surface-alt); padding:12px; border-radius:8px; }}
.error {{ background:#fff0f0; padding:12px; border-radius:8px; color:#8b2020; }}
.login {{ max-width:450px; margin:50px auto; }}
code {{ overflow-wrap:anywhere; }}
.theme-row {{ display:flex; align-items:center; justify-content:space-between; gap:15px; flex-wrap:wrap; }}
.theme-status {{ color:var(--muted); }}

@media(max-width:760px) {{
    .sidebar {{
        position: static;
        top: auto;
        left: auto;
        bottom: auto;
        width: auto;
        min-height: 0;
        overflow-y: visible;
        display: grid;
        grid-template-columns: repeat(2, minmax(0, 1fr));
    }}

    .side-label {{ grid-column: 1 / -1; }}
    .side-spacer {{ display: none; }}

    .app-shell {{
        display: block;
        grid-template-columns: 1fr;
        margin: 8px;
    }}

    .topbar {{ margin: 0 8px; }}
    main {{ padding-top: 0; }}
}}


/* Layout fix: allow all page content to remain accessible */
html, body {{
    min-height: 100%;
}}

body {{
    height: auto;
    min-height: 100vh;
    overflow-x: hidden;
    overflow-y: auto;
}}

.app-shell {{
    height: auto;
    min-height: 0;
    overflow: visible;
}}

main {{
    height: auto;
    min-height: 0;
    overflow: visible;
}}


/* Professional appearance upgrade */
:root {{
    --page: #f3f6fb;
    --surface: #ffffff;
    --surface-alt: #eaf0f8;
    --text: #172338;
    --muted: #64748b;
    --border: #dbe4f0;
    --navy: #163b65;
    --navy-hover: #24568a;
    --link: #245fa6;
    --sidebar: #ffffff;
    --shadow: 0 5px 18px rgba(25, 48, 80, 0.07);
}}

body {{
    font-family: "Segoe UI", Arial, sans-serif;
    font-size: 14px;
    line-height: 1.6;
    -webkit-font-smoothing: antialiased;
}}

.topbar {{
    min-height: 64px;
    padding: 12px 22px;
    border-radius: 0 0 14px 14px;
    box-shadow: 0 4px 14px rgba(15, 35, 65, 0.14);
}}

.brand {{
    font-size: 16px;
    letter-spacing: 0.2px;
}}

.brand small {{
    font-size: 11px;
    font-weight: normal;
    opacity: 0.85;
}}

.sidebar {{
    background: var(--surface);
    border-color: var(--border);
    box-shadow: var(--shadow);
    padding: 18px 12px;
}}

.side-label {{
    letter-spacing: 1px;
    font-size: 11px;
    color: var(--muted);
}}

.side-link {{
    min-height: 40px;
    padding: 10px 11px;
    border-radius: 8px;
    transition: background 0.18s ease, color 0.18s ease;
}}

.side-link:hover {{
    background: var(--surface-alt);
    color: var(--link);
}}

main h1 {{
    font-size: 27px;
    line-height: 1.3;
    letter-spacing: -0.5px;
    margin-bottom: 14px;
}}

h2 {{
    font-size: 19px;
    line-height: 1.4;
}}

section {{
    padding: 19px;
    margin: 12px 0;
    border-radius: 13px;
    border-color: var(--border);
    box-shadow: var(--shadow);
}}

.dashboard-widgets {{
    gap: 16px;
    margin: 16px 0;
}}

.stat-card {{
    min-height: 135px;
    transition: box-shadow 0.2s ease, transform 0.2s ease;
}}

.stat-card:hover {{
    box-shadow: 0 9px 24px rgba(25, 48, 80, 0.11);
    transform: translateY(-2px);
}}

.stat-card h2 {{
    font-size: 29px;
    font-weight: 700;
    letter-spacing: -0.5px;
}}

button,
.button,
.dashboard-button {{
    border-radius: 8px;
    font-weight: 600;
    transition: background 0.18s ease, box-shadow 0.18s ease;
}}

button:focus-visible,
.button:focus-visible,
.dashboard-button:focus-visible,
.side-link:focus-visible,
input:focus-visible,
select:focus-visible,
textarea:focus-visible {{
    outline: 3px solid #78aef5;
    outline-offset: 2px;
}}

input,
select,
textarea {{
    border-radius: 8px;
    padding: 11px 12px;
}}

input:focus,
select:focus,
textarea:focus {{
    border-color: #5289cb;
    box-shadow: 0 0 0 3px rgba(82, 137, 203, 0.13);
}}

.table {{
    border: 1px solid var(--border);
    border-radius: 10px;
    overflow-x: auto;
}}

table {{
    background: var(--surface);
}}

th {{
    font-weight: 700;
    color: var(--text);
}}

.note {{
    border: 1px solid var(--border);
    border-radius: 10px;
}}

@media (max-width: 600px) {{
    .topbar {{
        padding: 10px 12px;
    }}

    main h1 {{
        font-size: 23px;
    }}

    section {{
        padding: 14px;
    }}

    .stat-card h2 {{
        font-size: 25px;
    }}
}}


/* Dashboard analytics polish */
.dashboard-welcome {{
    color: var(--muted);
    font-size: 15px;
    max-width: 760px;
    margin: 0 0 22px;
}}

.dashboard-widgets {{
    align-items: stretch;
    gap: 18px;
}}

.stat-card {{
    position: relative;
    overflow: hidden;
    background: var(--surface);
    border: 1px solid var(--border);
    border-top: 3px solid #3978bd;
    border-radius: 12px;
    padding: 20px;
}}

.stat-card:nth-child(2) {{
    border-top-color: #159a80;
}}

.stat-card:nth-child(3) {{
    border-top-color: #805ad5;
}}

.stat-card:nth-child(4) {{
    border-top-color: #d69e2e;
}}

.stat-card .summary-label {{
    font-size: 12px;
    font-weight: 700;
    letter-spacing: 0.65px;
    text-transform: uppercase;
    margin-bottom: 12px;
}}

.stat-card h2 {{
    line-height: 1.3;
    margin: 0 0 12px;
    overflow-wrap: anywhere;
}}

.stat-card > span,
.stat-card > a {{
    font-size: 13px;
}}

.stat-card > span {{
    color: var(--muted);
}}

.stat-card > a {{
    display: inline-block;
    margin-top: 4px;
    font-weight: 600;
    text-decoration: none;
}}

.stat-card > a:hover {{
    text-decoration: underline;
}}

.chart-widget h2 {{
    margin-top: 0;
    margin-bottom: 4px;
}}

.chart-widget > p {{
    line-height: 1.5;
    margin-top: 4px;
}}

.chart-plot {{
    border-radius: 8px 8px 0 0;
}}

.subject-bar {{
    transition: opacity 0.2s ease;
}}

.subject-column:hover .subject-bar {{
    opacity: 0.8;
}}


/* Square Windows-style flip tiles */

/* Compact Windows-style flip tiles */
/* Two tiles per row */
.dashboard-widgets {{
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    align-items: stretch;
    margin: 12px 0;
}}

/* Each tile */
.dashboard-widgets {{
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 10px;
    align-items: stretch;
    margin: 10px 0;
}}

.dashboard-widgets > .flip-card {{
    width: 100%;
    height: 145px;
    min-height: 0;
    aspect-ratio: auto;
    perspective: 1000px;
    background: transparent;
    border: 0;
    padding: 0;
    margin: 0;
}}

.flip-card-inner {{
    position: relative;
    width: 100%;
    height: 100%;
    min-height: 0;
    transform-style: preserve-3d;
    transition: transform 0.6s ease;
}}

.flip-card-face {{
    position: absolute;
    inset: 0;
    box-sizing: border-box;
    width: 100%;
    height: 100%;
    padding: 12px;
    overflow: auto;
    border-radius: 12px;
    backface-visibility: hidden;
    -webkit-backface-visibility: hidden;
}}

.flip-card-front {{
    transform: rotateY(0deg);
}}

.flip-card-back {{
    transform: rotateY(180deg);
}}

.flip-card-inner {{
    transition: transform 0.5s;
    transform-style: preserve-3d;
}}

/* Do not flip automatically on hover or focus */
.flip-card:hover .flip-card-inner,
.flip-card:focus-within .flip-card-inner {{
    transform: none;
}}

/* Flip only when the widget is clicked */
.flip-card.is-flipped .flip-card-inner {{
    transform: rotateY(180deg);
}}

/* Keep the clicked widget flipped even while hovered */
.flip-card.is-flipped:hover .flip-card-inner,
.flip-card.is-flipped:focus-within .flip-card-inner {{
    transform: rotateY(180deg);
}}


/* Both faces occupy the same plane */
.flip-card-face {{
    position: absolute;
    inset: 0;
    box-sizing: border-box;
    width: 100%;
    height: 100%;
    backface-visibility: hidden;
    -webkit-backface-visibility: hidden;
    border-radius: 12px;
    padding: 14px;
    overflow: auto;
}}

.flip-card-front {{
    transform: rotateY(0deg);
}}

.flip-card-back {{
    transform: rotateY(180deg);
}}

.flip-card:nth-child(2) .flip-card-front {{ background: #147d70; }}
.flip-card:nth-child(3) .flip-card-front {{ background: #6246a5; }}
.flip-card:nth-child(4) .flip-card-front {{ background: #a66a12; }}

.flip-card-front .summary-label {{
    color: #ffffff;
    opacity: 0.9;
    font-size: 10px;
    font-weight: 700;
    letter-spacing: 0.6px;
    text-transform: uppercase;
    margin: 0;
}}

.flip-card-front h2 {{
    color: #ffffff;
    font-size: clamp(17px, 1.6vw, 25px);
    line-height: 1.2;
    margin: 0;
}}

.flip-card-front span,
.flip-card-front a {{
    color: #ffffff;
    font-size: 11px;
}}

.flip-card-front a {{
    font-weight: 700;
}}

.flip-card-back {{
    color: var(--text);
    background: var(--surface);
    border-top: 3px solid var(--navy);
    transform: rotateY(180deg);
}}

.flip-card-back h2 {{
    font-size: 15px;
    line-height: 1.3;
    margin: 0;
}}

.flip-card-back p {{
    color: var(--muted);
    font-size: 11px;
    line-height: 1.4;
    margin: 0;
}}

.flip-card-back a {{
    font-size: 11px;
    font-weight: 700;
    text-decoration: none;
}}

/* Make the graph compact so it fits beneath the tiles */
.chart-widget {{
    grid-column: 1 / -1;
    padding: 12px !important;
    margin: 0 !important;
}}

.chart-widget h2 {{
    font-size: 17px;
    margin: 0 0 3px;
}}

.chart-widget > p {{
    font-size: 12px;
    margin: 3px 0 8px;
}}

.chart-plot {{
    min-height: 130px;
    max-height: 160px;
}}

@media (max-width: 900px) {{
.dashboard-widgets {{
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    align-items: stretch;
    margin: 12px 0;
}}

@media (max-width: 600px) {{
    .dashboard-widgets {{
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 8px;
    }}

    .flip-card-face {{
        padding: 10px;
        gap: 6px;
    }}

    .flip-card-front h2 {{
        font-size: 17px;
    }}

    .flip-card-back h2 {{
        font-size: 13px;
    }}

    .flip-card-back p {{
        font-size: 10px;
    }}
}}

@media (prefers-reduced-motion: reduce) {{
    .flip-card-inner {{
        transition: none;
    }}
}}


@media (max-width: 600px) {{
    .dashboard-widgets {{
        grid-template-columns: minmax(0, 1fr);
    }}

    .flip-card {{
        aspect-ratio: 1 / 0.72;
        min-height: 250px;
    }}
}}

@media (prefers-reduced-motion: reduce) {{
    .flip-card-inner {{
        transition: none;
    }}
}}


/* Dashboard layout: two tiles across, two rows */
.dashboard-widgets {{
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 12px;
    align-items: stretch;
    margin: 12px 0;
}}

.dashboard-widgets > .flip-card {{
    aspect-ratio: 1 / 0.72;
    min-height: 0;
    height: auto;
}}

.dashboard-widgets > .chart-widget {{
    grid-column: 1 / -1;
    min-width: 0;
    padding: 12px;
    margin: 0;
}}

/* Align the percentage scale with the graph */
.subject-chart {{
    display: grid;
    grid-template-columns: 42px minmax(0, 1fr);
    gap: 8px;
    margin-top: 12px;
    align-items: stretch;
}}

.chart-y-axis {{
    height: 150px;
    min-height: 150px;
    padding-bottom: 27px;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    align-items: flex-end;
    font-size: 11px;
    color: var(--muted);
}}

.chart-main {{
    min-width: 0;
    overflow-x: auto;
}}

.chart-plot {{
    box-sizing: border-box;
    width: 100%;
    min-width: 360px;
    height: 150px;
    min-height: 150px;
    max-height: 150px;
    display: flex;
    align-items: stretch;
    gap: 10px;
    background:
        repeating-linear-gradient(
            to top,
            transparent 0,
            transparent calc(20% - 1px),
            var(--border) calc(20% - 1px),
            var(--border) 20%
        );
    background-size: 100% calc(100% - 27px);
    background-repeat: no-repeat;
    border-bottom: 1px solid var(--border);
}}

.subject-column {{
    flex: 1 1 0;
    min-width: 58px;
    height: 100%;
    position: relative;
}}

.subject-plot-column {{
    height: calc(100% - 27px);
    position: relative;
}}

.subject-bar {{
    position: absolute;
    bottom: 0;
    left: 25%;
    width: 50%;
    border-radius: 5px 5px 0 0;
}}

.subject-value {{
    position: absolute;
    left: 0;
    width: 100%;
    text-align: center;
    font-size: 11px;
    font-weight: bold;
    color: var(--text);
    white-space: nowrap;
}}

.subject-name {{
    height: 27px;
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
    text-align: center;
    font-size: 10px;
    color: var(--muted);
    padding-top: 5px;
}}

.chart-x-title,
.chart-y-title {{
    text-align: center;
    font-size: 11px;
    color: var(--muted);
    margin: 5px 0 0;
}}

/* Keep two tiles per row on smaller screens too */
@media (max-width: 600px) {{
    .dashboard-widgets {{
        grid-template-columns: repeat(2, minmax(0, 1fr));
        gap: 8px;
    }}

    .dashboard-widgets > .chart-widget {{
        grid-column: 1 / -1;
    }}

    .flip-card-face {{
        padding: 9px;
    }}

    .flip-card-front h2 {{
        font-size: 16px;
    }}
}}


/* Total learners: compact click-to-flip widget */
.click-flip-card:hover .flip-card-inner,
.click-flip-card:focus-within .flip-card-inner {{
    transform: rotateY(0deg);
}}

.click-flip-card.is-flipped .flip-card-inner {{
    transform: rotateY(180deg);
}}

.click-flip-card .flip-card-face {{
    display: flex;
    flex-direction: column;
    align-items: flex-start;
    gap: 6px;
    padding: 12px;
}}

.click-flip-card .flip-card-face h2 {{
    margin: 0;
    font-size: 24px;
}}

.click-flip-card .flip-card-face p {{
    margin: 3px 0;
}}

.click-flip-card .widget-flip-button {{
    margin-top: auto;
    padding: 7px 11px;
    border: 1px solid var(--border);
    border-radius: 7px;
    background: var(--surface, #ffffff);
    color: var(--text);
    cursor: pointer;
    font: inherit;
    font-size: 12px;
}}

.click-flip-card .widget-flip-button:hover {{
    border-color: #3b82f6;
}}

.click-flip-card .learner-performance-list {{
    list-style: none;
    padding: 0;
    margin: 0;
    width: 100%;
    display: flex;
    flex-direction: column;
    gap: 4px;
}}

.click-flip-card .learner-performance-row {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    gap: 8px;
    font-size: 12px;
}}

.click-flip-card .flip-card-back {{
    overflow-y: auto;
}}

.click-flip-card .flip-card-back a {{
    font-size: 12px;
}}



/* All dashboard tiles flip only when clicked */
.flip-card .flip-card-inner {{
    transition: transform 0.5s ease;
    transform-style: preserve-3d;
}}

.flip-card.is-flipped .flip-card-inner {{
    transform: rotateY(180deg);
}}

/* Do not flip automatically on hover */
.flip-card:hover .flip-card-inner {{
    transform: none;
}}

.flip-card.is-flipped:hover .flip-card-inner {{
    transform: rotateY(180deg);
}}

.flip-card {{
    cursor: pointer;
}}


</style>


<script>
(function() {{
    try {{
        if (localStorage.getItem('jss-theme') === 'dark') {{
            document.documentElement.classList.add('dark-start');
        }}
    }} catch (e) {{}}
}})();
</script>
</head>
<body>
<header class="topbar">
  <a class="brand" href="{destination if user else '/login'}"><span class="brand-mark">🎓</span><span>EduTrack Kenya<br><small>{esc(SCHOOL)} · Assessment Portal</small></span></a>
  <div class="top-actions">{('<a href="/settings" title="Settings">⚙ Settings</a><span title="Account">♙</span>' if user else '')}</div>
</header>
<div class="app-shell">
{sidebar}
<main>{body}</main>
</div>
<script>
(function() {{
  try {{
    if (localStorage.getItem('jss-theme') === 'dark') document.body.classList.add('dark-theme');
  }} catch (e) {{}}
  var toggle = document.getElementById('theme-toggle');
  
  var status = document.getElementById('theme-status');
  var dark = document.body.classList.contains('dark-theme');

  if (status) {{
    status.textContent = dark
      ? 'Dark theme is enabled.'
      : 'Light theme is enabled.';
  }}

  if (toggle) {{
    toggle.textContent = dark
      ? 'Switch to light theme'
      : 'Switch to dark theme';
  }}

  if (toggle) {{
    toggle.addEventListener('click', function() {{
      var dark = !document.body.classList.contains('dark-theme');
      document.body.classList.toggle('dark-theme', dark);
      try {{ localStorage.setItem('jss-theme', dark ? 'dark' : 'light'); }} catch (e) {{}}
      var status = document.getElementById('theme-status');
      if (status) status.textContent = dark ? 'Dark theme is enabled.' : 'Light theme is enabled.';
      toggle.textContent = dark ? 'Switch to light theme' : 'Switch to dark theme';
    }});
  }}
}})();
</script>
</body>
</html>"""
    return HTMLResponse(document, status_code=status_code)


def create_initial_teacher():
    email = os.getenv("TEACHER_EMAIL", "").strip().lower()
    password = os.getenv("TEACHER_PASSWORD", "")

    if not email or not password:
        print(
            "Set TEACHER_EMAIL and TEACHER_PASSWORD in .env "
            "to create the initial teacher account."
        )
        return

    if len(password) < 12:
        print("TEACHER_PASSWORD must be at least 12 characters.")
        return

    with DB() as db:
        existing = db.scalar(
            select(User).where(User.email == email)
        )


        if not existing:
            school = db.scalar(
                select(School).where(
                    School.name == SCHOOL
                )
            )

            if not school:
                print(
                    f"Cannot create initial teacher: "
                    f"school '{SCHOOL}' was not found."
                )
                return

            if school.status != "active":
                print(
                    f"Cannot create initial teacher: "
                    f"school '{SCHOOL}' is not active."
                )
                return

            db.add(
                User(
                    school_id=school.id,
                    email=email,
                    password_hash=passwords.hash(password),
                    role="teacher",
                )
            )
            db.commit()
            print(f"Initial teacher account created for {email}")


create_initial_teacher()


# ============================================================
# 4. LOGIN AND LOGOUT
# ============================================================

@app.get("/")
def root(request: Request):
    with DB() as db:
        user = current_user(request, db)

        if user:
            return go(
                "/teacher" if user.role == "teacher" else "/parent"
            )

    return go("/login")


# --------------------------------------------------
# SHARED LOGIN PAGE
# --------------------------------------------------

def render_login_page(
    request: Request,
    portal: str | None = None,
):
    if portal == "teacher":
        heading = "Teacher Portal Login"
        description = "Sign in to access the teacher dashboard."
        form_action = "/teacher/login"
    elif portal == "parent":
        heading = "Parent Portal Login"
        description = "Sign in to access the parent dashboard."
        form_action = "/parent/login"
    else:
        heading = "Sign in"
        description = f"Access the {esc(SCHOOL)} assessment portal."
        form_action = "/login"

    body = f"""
    <section class="login">
        <h1>{heading}</h1>
        <p>{description}</p>

        <form method="post" action="{form_action}">
            <input type="hidden" name="csrf_token"
                   value="{esc(csrf(request))}">

            <label>Email
                <input type="email" name="email" required
                       autocomplete="username">
            </label><br>

            <label>Password
                <input type="password" name="password" required
                       autocomplete="current-password">
            </label><br>

            <button type="submit">Sign in</button>
        </form>
    </section>
    """

    return page(request, heading, body)


# --------------------------------------------------
# LOGIN URLS
# --------------------------------------------------

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return render_login_page(request)


@app.get("/teacher/login", response_class=HTMLResponse)
def teacher_login_page(request: Request):
    return render_login_page(request, "teacher")


@app.get("/parent/login", response_class=HTMLResponse)
def parent_login_page(request: Request):
    return render_login_page(request, "parent")


# --------------------------------------------------
# SHARED LOGIN PROCESS
# --------------------------------------------------

def process_login(
    request: Request,
    email: str,
    password: str,
    csrf_token: str,
    expected_role: str | None = None,
):
    check_csrf(request, csrf_token)

    with DB() as db:
        user = db.scalar(
            select(User).where(
                User.email == email.strip().lower()
            )
        )

        # Reject invalid credentials.
        if not user or not passwords.verify(
            password, user.password_hash
        ):
            retry_url = (
                "/teacher/login"
                if expected_role == "teacher"
                else "/parent/login"
                if expected_role == "parent"
                else "/login"
            )

            body = f"""
            <section class="login">
                <h1>Sign in failed</h1>
                <p>Check your email and password.</p>
                <a class="button" href="{retry_url}">Try again</a>
            </section>
            """

            return page(
                request,
                "Sign in failed",
                body,
                status_code=401,
            )

        # Ensure the account belongs to the selected portal.
        if expected_role and user.role != expected_role:
            body = """
            <section class="login">
                <h1>Wrong portal</h1>
                <p>
                    This account does not have access to this portal.
                    Please use the correct login address.
                </p>
                <p>
                    <a href="/teacher/login">Teacher login</a>
                    |
                    <a href="/parent/login">Parent login</a>
                </p>
            </section>
            """

            return page(
                request,
                "Wrong portal",
                body,
                status_code=403,
            )

        # Start a fresh authenticated session.
        request.session.clear()
        request.session["user_id"] = user.id
        request.session["csrf_token"] = secrets.token_urlsafe(32)

        return go(
            "/teacher" if user.role == "teacher" else "/parent"
        )


# --------------------------------------------------
# LOGIN FORM SUBMISSIONS
# --------------------------------------------------

@app.post("/login")
def login(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(...),
):
    return process_login(
        request, email, password, csrf_token
    )


@app.post("/teacher/login")
def teacher_login(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(...),
):
    return process_login(
        request,
        email,
        password,
        csrf_token,
        expected_role="teacher",
    )


@app.post("/parent/login")
def parent_login(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(...),
):
    return process_login(
        request,
        email,
        password,
        csrf_token,
        expected_role="parent",
    )


# --------------------------------------------------
# LOGOUT
# --------------------------------------------------

@app.post("/logout")
def logout(
    request: Request,
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)
    request.session.clear()
    return go("/login")

# ============================================================
# 5. TEACHER DASHBOARD
# ============================================================

@app.get("/teacher", response_class=HTMLResponse)
def teacher_dashboard(request: Request):
    with DB() as db:
        user = require_role(request, db, "teacher")

        # Load only learners belonging to this teacher's school.
        students = db.scalars(
            select(Student)
            .where(Student.school_id == user.school_id)
            .order_by(Student.grade, Student.name)
        ).all()

        # Get the search term from the URL, if provided.
        search_query = request.query_params.get("search", "").strip()

        # Start with all learners, then filter if a search was entered.
        matching_students = students

        if search_query:
            query_lower = search_query.casefold()
            matching_students = [
                s for s in students
                if query_lower in (s.name or "").casefold()
                or query_lower in (s.admission or "").casefold()
            ]

        parents = db.scalars(
            select(User)
            .where(User.role == "parent")
            .where(User.school_id == user.school_id)
            .order_by(User.email)
        ).all()

        results = db.scalars(
            select(Result)
            .join(Student, Result.student_id == Student.id)
            .where(Student.school_id == user.school_id)
            .order_by(Result.id.desc())
            .limit(50)
        ).all()

        # Load every assessment for accurate dashboard statistics.
        all_results = db.scalars(
            select(Result)
            .join(Student, Result.student_id == Student.id)
            .where(Student.school_id == user.school_id)
        ).all()

        subject_scores = {}
        student_scores = {}
        all_percentages = []

        for result in all_results:
            if result.maximum <= 0:
                continue

            percentage = result.marks / result.maximum * 100
            all_percentages.append(percentage)

            subject_scores.setdefault(result.subject, []).append(
                percentage
            )

            student_scores.setdefault(
                result.student_id,
                {
                    "name": result.student.name,
                    "scores": [],
                },
            )["scores"].append(percentage)

        subject_averages = {
            subject: sum(scores) / len(scores)
            for subject, scores in subject_scores.items()
            if scores
        }

        overall_average = (
            sum(all_percentages) / len(all_percentages)
            if all_percentages else 0
        )
        best_subject = (
            max(subject_averages, key=subject_averages.get)
            if subject_averages else "No results yet"
        )

        best_student = (
            max(
                student_scores.values(),
                key=lambda item: (
                    sum(item["scores"]) / len(item["scores"])
                ),
            )
            if student_scores else None
        )

        best_student_average = (
            sum(best_student["scores"]) / len(best_student["scores"])
            if best_student else 0
        )

        # Count assessed learners and calculate each learner's average.
        learner_scores = {}

        for result in all_results:
            if result.maximum and result.maximum > 0:
                learner_id = result.student.id
                percentage = (result.marks / result.maximum) * 100
                learner_scores.setdefault(learner_id, []).append(percentage)

        assessed_learners = len(learner_scores)

        learner_averages = [
            sum(scores) / len(scores)
            for scores in learner_scores.values()
            if scores
        ]

        # Starter performance bands; these can be customized later.
        performance_levels = {
            "High (80–100%)": 0,
            "On track (60–79.9%)": 0,
            "Developing (40–59.9%)": 0,
            "Needs support (below 40%)": 0,
        }

        for average in learner_averages:
            if average >= 80:
                performance_levels["High (80–100%)"] += 1
            elif average >= 60:
                performance_levels["On track (60–79.9%)"] += 1
            elif average >= 40:
                performance_levels["Developing (40–59.9%)"] += 1
            else:
                performance_levels["Needs support (below 40%)"] += 1

        performance_rows = "".join(
            f'<li class="learner-performance-row">'
            f'<span>{label}</span>'
            f'<strong>{count}</strong>'
            f'</li>'
            for label, count in performance_levels.items()
        )

        chart_rows = "".join(
            f'<div class="subject-column">'
            f'<div class="subject-plot-column">'
            f'<span class="subject-value" '
            f'style="bottom:calc({score:.2f}% + 4px)">'
            f'{score:.1f}%</span>'
            f'<div class="subject-bar" '
            f'style="height:{score:.2f}%"></div>'
            f'</div>'
            f'<div class="subject-name">{esc(subject)}</div>'
            f'</div>'
            for subject, score in subject_averages.items()
        ) or '<p>No assessment data available yet.</p>'

        student_options = "".join(
            f'<option value="{s.id}">{esc(s.name)} — '
            f'Grade {s.grade} — '
            f'{esc(s.admission or "No admission number")}</option>'
            for s in students
        )

        parent_options = "".join(
            f'<option value="{p.id}">{esc(p.email)}</option>'
            for p in parents
        )

        subject_options = "".join(
            f'<option value="{esc(s)}">{esc(s)}</option>'
            for s in SUBJECTS
        )

        grade_options = "".join(
            f'<option value="{g}">Grade {g}</option>'
            for g in GRADES
        )

        learner_rows = "".join(
            f"<tr>"
            f"<td>{esc(s.admission or '—')}</td>"
            f'<td><a href="/teacher/learner/{s.id}">{esc(s.name)}</a></td>'
            f"<td>{s.grade}</td>"
            f"<td>{len(s.results)}</td>"
            f"</tr>"
            for s in students
        ) or (
            '<tr><td colspan="4">No matching learners found.</td></tr>'
        ) or '<tr><td colspan="4">No learners registered.</td></tr>'

        result_rows = "".join(
            f"<tr><td>{esc(r.student.name)}</td>"
            f"<td>{esc(r.subject)}</td>"
            f"<td>{esc(r.assessment)}</td>"
            f"<td>{r.marks:g}/{r.maximum:g}</td>"
            f"<td>{r.result_date}</td></tr>"
            for r in results
        ) or '<tr><td colspan="5">No results recorded.</td></tr>'

        body = f"""
        <h1>Hello, {esc(user.name or 'Teacher')}!</h1>
        <p class="dashboard-welcome">
            Welcome to your school management workspace.
            Manage learners, parents and academic results from one place.
        </p>


        <div class="dashboard-widgets">
           

<section class="stat-card flip-card click-flip-card total-learners-widget"
         tabindex="0"
         aria-label="Total learners and assessment overview"
         onclick="this.classList.toggle('is-flipped')">

    <div class="flip-card-inner">

        <div class="flip-card-face flip-card-front">
            <p class="summary-label">Total learners</p>

            <h2>{len(students)}</h2>

            <p>Registered learners</p>

            <button
                type="button"
                class="widget-flip-button"
                onclick="this.closest('.click-flip-card').classList.toggle('is-flipped')">
            
<a
    href="/teacher/analysis"
    class="widget-flip-button"
    onclick="event.stopPropagation()">
    View analysis →
</a>

        </div>

        <div class="flip-card-face flip-card-back">
            <h2>Learner overview</h2>

            <p>
                <strong>{len(students)}</strong> learners are registered
                at this school.
            </p>

            <p>
                <strong>{assessed_learners}</strong> learners have at least
                one valid recorded assessment result.
            </p>

            <h3>Performance levels</h3>

            <ul class="learner-performance-list">
                {performance_rows}
            </ul>

            <a href="/teacher/learners">
                View all learners →
            </a>

            <button
                type="button"
                class="widget-flip-button"
                onclick="this.closest('.click-flip-card').classList.toggle('is-flipped')">
                ← Back
            </button>
        </div>

    </div>
</section>

<section class="stat-card flip-card"
         tabindex="0"
         onclick="this.classList.toggle('is-flipped')">
    <div class="flip-card-inner">
        <div class="flip-card-face flip-card-front">
            <p class="summary-label">Best subject</p>
            <h2>{esc(best_subject)}</h2>
            <span>
                {f"{subject_averages[best_subject]:.1f}% average"
                 if best_subject in subject_averages
                 else "No marks recorded"}
            </span>
        </div>
        <div class="flip-card-face flip-card-back">
            <h2>Subject analysis</h2>
            <p>
                {(
                    f"{best_subject} currently has the highest recorded "
                    f"subject average of {subject_averages[best_subject]:.1f}%."
                )
                if best_subject in subject_averages
                else "Subject comparisons will appear when assessment marks are recorded."}
            </p>
        </div>
    </div>
</section>

<section class="stat-card flip-card"
         tabindex="0"
         onclick="this.classList.toggle('is-flipped')">
    <div class="flip-card-inner">
        <div class="flip-card-face flip-card-front">
            <p class="summary-label">Overall average</p>
            <h2>{overall_average:.1f}%</h2>
            <span>All recorded assessments</span>
        </div>
        <div class="flip-card-face flip-card-back">
            <h2>Performance overview</h2>
            <p>
                This average is calculated from valid recorded assessment
                percentages across the school.
            </p>
            <p>
                Assessments included: {len(all_percentages)}
            </p>
        </div>
    </div>
</section>

<section class="stat-card flip-card"
         tabindex="0"
         onclick="this.classList.toggle('is-flipped')">
    <div class="flip-card-inner">
        <div class="flip-card-face flip-card-front">
            <p class="summary-label">Best student</p>
            <h2>
                {esc(best_student["name"] or "No results yet")
                 if best_student else "No results yet"}
            </h2>
            <span>
                {f"{best_student_average:.1f}% average"
                 if best_student else "No marks recorded"}
            </span>
        </div>
        <div class="flip-card-face flip-card-back">
            <h2>Learner analysis</h2>
            <p>
                {(
                    f"{esc(best_student['name'] or 'This learner')} has "
                    f"the highest average among learners with recorded "
                    f"results: {best_student_average:.1f}%."
                )
                if best_student else
                "Learner performance analysis will appear when assessment marks are available."}
            </p>
        </div>
    </div>
</section>


            <section class="chart-widget">
                <h2>Subject performance</h2>
                <p>Average percentage by subject</p>

                <div class="subject-chart">
                    <div class="chart-y-axis">
                        <span>100%</span>
                        <span>80%</span>
                        <span>60%</span>
                        <span>40%</span>
                        <span>20%</span>
                        <span>0%</span>
                    </div>

                    <div class="chart-main">
                        <div class="chart-plot">
                            {chart_rows}
                        </div>
                        <p class="chart-x-title">Subjects</p>
                    </div>
                </div>

                <p class="chart-y-title">Average marks (%)</p>
            </section>
        </div>
        """

        return page(request, "Teacher dashboard", body, user)


# ============================================================
# 5a. ADD LEARNER AND SCHOOL MANAGEMENT PAGE
# ============================================================

@app.get("/teacher/add-learner", response_class=HTMLResponse)
def teacher_add_learner(request: Request):
    with DB() as db:
        user = require_role(request, db, "teacher")

        students = db.scalars(
         select(Student)
         .where(Student.school_id == user.school_id)
         .order_by(Student.grade, Student.name)
        ).all()

        parents = db.scalars(
            select(User)
            .where(User.role == "parent")
            .where(User.school_id == user.school_id)
            .order_by(User.email)
        ).all()

        student_options = "".join(
            f'<option value="{s.id}">{esc(s.name)} — '
            f'Grade {s.grade} — '
            f'{esc(s.admission or "No admission number")}</option>'
            for s in students
        )

        parent_options = "".join(
            f'<option value="{p.id}">{esc(p.email)}</option>'
            for p in parents
        )

        subject_options = "".join(
            f'<option value="{esc(subject)}">{esc(subject)}</option>'
            for subject in SUBJECTS
        )

        grade_options = "".join(
            f'<option value="{grade}">Grade {grade}</option>'
            for grade in GRADES
        )

       
        learner_rows = "".join(
            f"<tr>"
            f"<td>{esc(s.admission or '—')}</td>"
            f'<td><a href="/teacher/learner/{s.id}">{esc(s.name)}</a></td>'
            f"<td>{s.grade}</td>"
            f"<td>{len(s.results)}</td>"
            f"</tr>"
            for s in students
        ) or (
            '<tr><td colspan="4">No learners registered.</td></tr>'
        )


        body = f"""
        <h1>Add Learner</h1>
        <p>
            Register learners and manage parent accounts, parent links,
            and assessment marks from this page.
        </p>

        <section>
            <h2>Register learner</h2>
            <form class="grid" method="post" action="/students">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Full name
                    <input name="name" maxlength="160" required>
                </label>

                <label>Grade
                    <select name="grade" required>
                        {grade_options}
                    </select>
                </label>

                <label>Admission number
                    <input name="admission" maxlength="80">
                </label>

                <button type="submit">Save learner</button>
            </form>
        </section>

        <section>
            <h2>Create parent account</h2>
            <p class="note">
                The temporary password must be at least 12 characters.
            </p>

            <form class="grid" method="post" action="/parents">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Parent email
                    <input type="email" name="email" required>
                </label>

                <label>Temporary password
                    <input type="password" name="password"
                           minlength="12" required>
                </label>

                <button type="submit">Create parent account</button>
            </form>
        </section>

        <section>
            <h2>Reset parent password</h2>
            <p>
                Generate a new temporary password for an existing parent.
            </p>

            <form class="grid" method="post"
                  action="/parents/reset-password">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Select parent
                    <select name="parent_id" required>
                        {parent_options}
                    </select>
                </label>

                <button type="submit">Generate new password</button>
            </form>
        </section>

        <section>
            <h2>Link parent to learner</h2>
            <form class="grid" method="post" action="/link">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Parent
                    <select name="parent_id" required>
                        {parent_options}
                    </select>
                </label>

                <label>Learner
                    <select name="student_id" required>
                        {student_options}
                    </select>
                </label>

                <button type="submit">Link parent</button>
            </form>
        </section>

        <section>
            <h2>Record assessment marks</h2>
            <form class="grid" method="post" action="/results">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Learner
                    <select name="student_id" required>
                        {student_options}
                    </select>
                </label>

                <label>Subject
                    <select name="subject" required>
                        {subject_options}
                    </select>
                </label>

                <label>Assessment
                    <input name="assessment" maxlength="120"
                           placeholder="End Term 1" required>
                </label>

                <label>Date
                    <input type="date" name="result_date"
                           value="{date.today().isoformat()}" required>
                </label>

                <label>Marks obtained
                    <input type="number" name="marks"
                           min="0" step="0.01" required>
                </label>

                <label>Maximum marks
                    <input type="number" name="maximum"
                           min="0.01" step="0.01"
                           value="100" required>
                </label>

                <button type="submit">Save assessment marks</button>
            </form>
        </section>

        <section>
            <h2>Learners ({len(students)})</h2>
            <p>
                <a class="button" href="/teacher/learners">
                    View all learners
                </a>
            </p>

            <label for="learner-search">Search learners</label>
            <input type="search" id="learner-search"
                   placeholder="Type a learner's name or admission number..."
                   autocomplete="off">

            <p id="learner-search-empty" hidden>
                No matching learners found.
            </p>

            <div class="table" id="learner-search-results" hidden>
                <table>
                    <thead>
                        <tr>
                            <th>Admission</th>
                            <th>Name</th>
                            <th>Grade</th>
                            <th>Results</th>
                        </tr>
                    </thead>
                    <tbody>
                        {learner_rows}
                    </tbody>
                </table>
            </div>
        </section>

        <script>
        (function() {{
            var input = document.getElementById("learner-search");
            var box = document.getElementById("learner-search-results");
            var empty = document.getElementById("learner-search-empty");
            var rows = box.querySelectorAll("tbody tr");

            input.addEventListener("input", function() {{
                var query = input.value.trim().toLowerCase();
                var shown = 0;

                rows.forEach(function(row) {{
                    var matches = query !== "" &&
                        row.textContent.toLowerCase().includes(query) &&
                        shown < 3;

                    row.hidden = !matches;

                    if (matches) {{
                        shown++;
                    }}
                }});

                box.hidden = query === "" || shown === 0;
                empty.hidden = query === "" || shown > 0;
            }});
        }})();
        </script>
        """

        return page(request, "Add Learner", body, user)

        

        

# ============================================================
# 5b. DISPLAY LEARNER RECORDS
# ============================================================

@app.get("/teacher/learners", response_class=HTMLResponse)
def teacher_all_learners(request: Request):
    with DB() as db:
        user = require_role(request, db, "teacher")

        # Load learners in alphabetical order by grade and name.
        

        
        students = db.scalars(
            select(Student)
            .where(Student.school_id == user.school_id)
            .order_by(Student.grade, Student.name)
        ).all()

        # Get the search term from the URL.
        search_query = request.query_params.get("search", "").strip()

        # By default, include all learners from this school.
        matching_students = students

        # Filter the learners when a search term is provided.
        if search_query:
            query_lower = search_query.casefold()
            matching_students = [
                s for s in students
                if query_lower in (s.name or "").casefold()
                or query_lower in (s.admission or "").casefold()
            ]



        # Show 10 learners by default, or 10 search matches.
        page_number = max(1, int(request.query_params.get("page", "1")))
        per_page = 10

        total_pages = max(1, (len(matching_students) + per_page - 1) // per_page)
        page_number = min(page_number, total_pages)

        start = (page_number - 1) * per_page
        end = start + per_page

        students = matching_students[start:end]

        learner_rows = "".join(
            f"<tr>"
            f"<td>{esc(s.admission or '—')}</td>"
            f'<td><a href="/teacher/learner/{s.id}">{esc(s.name)}</a></td>'
            f"<td>{s.grade}</td>"
            f"<td>{len(s.results)}</td>"
            f"</tr>"
            for s in students
        ) or (
            '<tr><td colspan="4">No matching learners found.</td></tr>'
        )

        total_learners = len(
            db.scalars(
                select(Student).where(
                    Student.school_id == user.school_id
                )
            ).all()
        )
        matching_count = len(matching_students)

        if search_query:
            list_message = (
                f"Found {matching_count} matching learner(s). "
                f"Showing {len(students)}."
            )
        else:
            list_message = (
                f"Showing {len(students)} of "
                f"{total_learners} registered learners."
            )

        body = f"""
        <h1>All Learners</h1>

        <section>
            <h2>Search learners</h2>


<form method="get" action="/teacher/learners"
      id="learner-search-form"
      style="display:flex; align-items:end; gap:8px; flex-wrap:wrap; margin-bottom:15px;">

    <label for="learner-search"
           style="width:280px; max-width:100%;">
        Learner name or admission number
        <input
            type="search"
            id="learner-search"
            name="search"
            value="{esc(search_query)}"
            placeholder="Search learners..."
            autocomplete="off"
            style="width:100%; box-sizing:border-box;"
        >
    </label>

    <button type="submit">Search</button>

    <a class="button" href="/teacher/learners">Reset</a>
</form>


            <div style="display:flex; justify-content:space-between; align-items:center; gap:12px; flex-wrap:wrap; margin:15px 0;">
    <p style="margin:0;">{list_message} Page {page_number} of {total_pages}</p>

    <div style="display:flex; gap:8px;">
        <a class="button"
           href="/teacher/learners?page={page_number - 1}&search={esc(search_query)}"
           style="pointer-events:{'none' if page_number == 1 else 'auto'}; opacity:{'0.5' if page_number == 1 else '1'};">
            Previous
        </a>

        <a class="button"
           href="/teacher/learners?page={page_number + 1}&search={esc(search_query)}"
           style="pointer-events:{'none' if page_number == total_pages else 'auto'}; opacity:{'0.5' if page_number == total_pages else '1'};">
            Next →
        </a>
    </div>
</div>

            <div style="width:100%; overflow-x:auto;">
                 <table style="width:100%; table-layout:auto; border-collapse:collapse; font-size:14px;">
                    <thead>
                        <tr>
                            <th>Admission</th>
                            <th>Name</th>
                            <th>Grade</th>
                            <th>Results</th>
                        </tr>
                    </thead>
                    <tbody>
                        {learner_rows}
                    </tbody>
                </table>
            </div>
        </section>




<script>
(function() {{
    const input = document.getElementById("learner-search");
    const form = document.getElementById("learner-search-form");

    if (!input || !form) return;

    input.addEventListener("keydown", function(event) {{
        if (event.key === "Enter") {{
            event.preventDefault();
            form.requestSubmit();
        }}
    }});
}})();
</script>

"""

    return page(request, "All Learners", body, user)



@app.get("/teacher/analysis", response_class=HTMLResponse)
def teacher_analysis(request: Request):
    with DB() as db:
        user = require_role(request, db, "teacher")

        students = db.scalars(
            select(Student)
            .where(Student.school_id == user.school_id)
            .order_by(Student.grade, Student.name)
        ).all()

        results = db.scalars(
            select(Result)
            .join(Student, Result.student_id == Student.id)
            .where(Student.school_id == user.school_id)
            .where(Result.maximum > 0)
        ).all()


        # Calculate each learner's average percentage.
        learner_scores = {}

        # Calculate average percentage for each subject.
        subject_scores = {}

        for result in results:
            percentage = result.marks / result.maximum * 100

            learner_scores.setdefault(
                result.student_id, []
            ).append(percentage)

            subject_scores.setdefault(
                result.subject, []
            ).append(percentage)

        learner_averages = {}

        for student in students:
            scores = learner_scores.get(student.id, [])
            if scores:
                learner_averages[student.id] = (
                    sum(scores) / len(scores)
                )

        # Sort learners by average percentage.
        ranked_students = sorted(
            [
                (student, learner_averages[student.id])
                for student in students
                if student.id in learner_averages
            ],
            key=lambda item: item[1],
            reverse=True,
        )

        best_rows = "".join(
            f"<tr>"
            f"<td>{esc(student.name)}</td>"
            f"<td>Grade {student.grade}</td>"
            f"<td>{average:.1f}%</td>"
            f"</tr>"
            for student, average in ranked_students[:10]
        ) or '<tr><td colspan="3">No assessment results yet.</td></tr>'

        # Rank all assessed learners from highest to lowest average.
        ranking_rows = "".join(
            f"<tr>"
            f"<td>{position}</td>"
            f"<td>{esc(student.name)}</td>"
            f"<td>Grade {student.grade}</td>"
            f"<td>{average:.1f}%</td>"
            f"<td>{'High' if average >= 80 else 'On track' if average >= 60 else 'Developing' if average >= 40 else 'Needs support'}</td>"
            f"</tr>"
            for position, (student, average) in enumerate(
                ranked_students, start=1
            )
        ) or '<tr><td colspan="5">No assessment results available yet.</td></tr>' 

        # Starter support threshold: below 40%.
        support_students = [
            (student, average)
            for student, average in ranked_students
            if average < 40
        ]

        support_rows = "".join(
            f"<tr>"
            f"<td>{esc(student.name)}</td>"
            f"<td>Grade {student.grade}</td>"
            f"<td>{average:.1f}%</td>"
            f"<td>Review assessment gaps and provide support.</td>"
            f"</tr>"
            for student, average in support_students
        ) or '<tr><td colspan="4">No learners currently fall below 40%.</td></tr>'

        # Learners who have not yet been assessed.
        unassessed = [
            student for student in students
            if student.id not in learner_averages
        ]

        unassessed_rows = "".join(
            f"<tr>"
            f"<td>{esc(student.name)}</td>"
            f"<td>Grade {student.grade}</td>"
            f"</tr>"
            for student in unassessed
        ) or '<tr><td colspan="2">All registered learners have assessment results.</td></tr>'

        subject_averages = {
            subject: sum(scores) / len(scores)
            for subject, scores in subject_scores.items()
            if scores
        }

        subject_rows = "".join(
            f"<tr>"
            f"<td>{esc(subject)}</td>"
            f"<td>{average:.1f}%</td>"
            f"</tr>"
            for subject, average in sorted(
                subject_averages.items(),
                key=lambda item: item[1],
                reverse=True,
            )
        ) or '<tr><td colspan="2">No subject results available yet.</td></tr>'

        # Editable starter performance bands.
        bands = {
            "High: 80–100%": 0,
            "On track: 60–79.9%": 0,
            "Developing: 40–59.9%": 0,
            "Needs support: below 40%": 0,
        }

        for average in learner_averages.values():
            if average >= 80:
                bands["High: 80–100%"] += 1
            elif average >= 60:
                bands["On track: 60–79.9%"] += 1
            elif average >= 40:
                bands["Developing: 40–59.9%"] += 1
            else:
                bands["Needs support: below 40%"] += 1

        band_rows = "".join(
            f"<tr><td>{label}</td><td>{count}</td></tr>"
            for label, count in bands.items()
        )

        overall_average = (
            sum(learner_averages.values())
            / len(learner_averages)
            if learner_averages else 0
        )

        body = f"""
        <h1>Learner Performance Analysis</h1>
        <p class="muted">
            Analysis based on recorded assessment percentages.
            A learner's average is calculated from their available results.
        </p>

        <p>
            <a class="dashboard-button" href="/teacher">
                ← Back to Teacher Dashboard
            </a>
        </p>

        <div class="stat-grid">
            <div class="stat-card">
                <div class="stat-label">Registered learners</div>
                <span class="stat-value">{len(students)}</span>
            </div>
            <div class="stat-card">
                <div class="stat-label">Learners assessed</div>
                <span class="stat-value">{len(learner_averages)}</span>
            </div>
            <div class="stat-card">
                <div class="stat-label">Overall learner average</div>
                <span class="stat-value">{overall_average:.1f}%</span>
            </div>
            <div class="stat-card">
                <div class="stat-label">Need extra support</div>
                <span class="stat-value">{len(support_students)}</span>
            </div>
        </div>

        <section>
            <h2>Subject performance</h2>
            <p class="muted">
                Subjects are ranked by their average percentage.
            </p>
            <div class="table">
                <table>
                    <thead>
                        <tr><th>Subject</th><th>Average</th></tr>
                    </thead>
                    <tbody>{subject_rows}</tbody>
                </table>
            </div>
        </section>

                <section>
            <h2>Complete Learner Ranking</h2>
            <p class="muted">
                All assessed learners ranked by their average assessment
                percentage, from highest to lowest.
            </p>
            <div class="table">
                <table>
                    <thead>
                        <tr>
                            <th>Rank</th>
                            <th>Learner</th>
                            <th>Grade</th>
                            <th>Average Marks</th>
                            <th>Performance</th>
                        </tr>
                    </thead>
                    <tbody>{ranking_rows}</tbody>
                </table>
            </div>
        </section>
        <section>
            <h2>Top-performing learners</h2>
            <p class="muted">
                The list shows the ten highest learner averages across
                their available assessment results.
            </p>
            <div class="table">
                <table>
                    <thead>
                        <tr>
                            <th>Learner</th><th>Grade</th><th>Average</th>
                        </tr>
                    </thead>
                    <tbody>{best_rows}</tbody>
                </table>
            </div>
        </section>

        <section>
            <h2>Learners needing additional support</h2>
            <p class="muted">
                Starter rule: average below 40%. Review individual subjects
                and assessments before deciding what intervention is needed.
            </p>
            <div class="table">
                <table>
                    <thead>
                        <tr>
                            <th>Learner</th><th>Grade</th>
                            <th>Average</th><th>Suggested action</th>
                        </tr>
                    </thead>
                    <tbody>{support_rows}</tbody>
                </table>
            </div>
        </section>

        <section>
            <h2>Performance distribution</h2>
            <p class="muted">
                These are configurable percentage bands, not a claim
                that every school level uses identical grade boundaries.
            </p>
            <div class="table">
                <table>
                    <thead>
                        <tr><th>Performance band</th><th>Learners</th></tr>
                    </thead>
                    <tbody>{band_rows}</tbody>
                </table>
            </div>
        </section>

        <section>
            <h2>Learners without recorded assessment results</h2>
            <div class="table">
                <table>
                    <thead>
                        <tr><th>Learner</th><th>Grade</th></tr>
                    </thead>
                    <tbody>{unassessed_rows}</tbody>
                </table>
            </div>
        </section>
        """

        return page(
            request,
            "Learner Performance Analysis",
            body,
            user=user,
        )



@app.post("/teacher/learner/{student_id}/results/{result_id}/edit")
def edit_learner_result(
    request: Request,
    student_id: int,
    result_id: int,
    subject: str = Form(...),
    assessment: str = Form(...),
    result_date: str = Form(...),
    marks: float = Form(...),
    maximum: float = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    subject = subject.strip()
    assessment = assessment.strip()

    if not subject or len(subject) > 120:
        raise HTTPException(status_code=400, detail="Invalid subject.")

    if not assessment or len(assessment) > 120:
        raise HTTPException(status_code=400, detail="Invalid assessment name.")


    if not (
        math.isfinite(marks)
        and math.isfinite(maximum)
        and maximum > 0
        and 0 <= marks <= maximum
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Invalid marks. The mark must be between "
                "zero and a positive maximum."
            ),
        )

    try:
        parsed_date = date.fromisoformat(result_date)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid assessment date.")

    with DB() as db:
        user = require_role(request, db, "teacher")

        student = db.scalar(
            select(Student).where(
                Student.id == student_id,
                Student.school_id == user.school_id,
            )
        )

        result = db.scalar(
            select(Result).where(
                Result.id == result_id,
                Result.student_id == student_id,
                Result.student.has(Student.school_id == user.school_id),
            )
        )

        if not student or not result:
            raise HTTPException(
                status_code=404,
                detail="Learner or assessment not found.",
            )
        result.subject = subject
        result.assessment = assessment
        result.result_date = parsed_date
        result.marks = marks
        result.maximum = maximum

        db.commit()

    return go(f"/teacher/learner/{student_id}")




@app.get("/teacher/learner/{student_id}", response_class=HTMLResponse)
def teacher_learner_profile(request: Request, student_id: int):
    with DB() as db:
        user = require_role(request, db, "teacher")

        student = db.scalar(
            select(Student).where(
                Student.id == student_id,
                Student.school_id == user.school_id,
            )
        )

        if not student:
            raise HTTPException(
                status_code=404,
                detail="Learner not found.",
            )
        results = sorted(
            student.results,
            key=lambda r: (r.result_date, r.subject),
            reverse=True,
        )

        
        # Prepare assessment results in chronological order for the graph.
        chart_results = sorted(
            student.results,
            key=lambda r: (r.result_date, r.id),
        )

        graph_points = []
        graph_circles = []
        graph_dates = []

        for i, r in enumerate(chart_results):
            score = (r.marks / r.maximum * 100) if r.maximum else 0
            x = 70 + i * (650 / max(1, len(chart_results) - 1))
            y = 225 - score * 1.9

            graph_points.append(f"{x:.1f},{y:.1f}")

            tooltip = (
                f"{esc(r.assessment)}: {score:.1f}% "
                f"({r.result_date.strftime('%d %b %Y')})"
            )

            graph_circles.append(
                f'<circle cx="{x:.1f}" cy="{y:.1f}" r="5" '
                f'fill="currentColor"><title>{tooltip}</title></circle>'
            )

            graph_dates.append(
                f'<text x="{x:.1f}" y="250" font-size="10" '
                f'text-anchor="end" '
                f'transform="rotate(-35 {x:.1f} 250)">'
                f'{r.result_date.strftime("%d %b %y")}</text>'
            )

        graph_grid = "".join(
            f'<line x1="70" y1="{225 - pct * 1.9}" '
            f'x2="730" y2="{225 - pct * 1.9}" '
            f'stroke="var(--border)" stroke-width="1"/>'
            f'<text x="58" y="{230 - pct * 1.9}" '
            f'text-anchor="end" font-size="11">{pct}%</text>'
            for pct in (0, 25, 50, 75, 100)
        )

        if chart_results:
            graph_svg = f"""
            <div style="width:100%; overflow-x:auto;">
                <svg viewBox="0 0 760 310"
                     role="img"
                     aria-label="Learner marks over time"
                     style="width:100%; min-width:420px;
                            max-height:360px; color:var(--text);">
                    {graph_grid}
                    <line x1="70" y1="35" x2="70" y2="225"
                          stroke="currentColor"/>
                    <line x1="70" y1="225" x2="730" y2="225"
                          stroke="currentColor"/>
                    <polyline points="{' '.join(graph_points)}"
                              fill="none" stroke="currentColor"
                              stroke-width="2.5"
                              stroke-linejoin="round"
                              stroke-linecap="round"/>
                    {''.join(graph_circles)}
                    {''.join(graph_dates)}
                    <text x="400" y="305" text-anchor="middle"
                          font-size="12">Assessment date</text>
                </svg>
            </div>
            """
        else:
            graph_svg = "<p>No assessment results available to graph yet.</p>"

        # Build the assessment table with an edit form for each result.
        result_rows = "".join(
            f"""
            <tr>
                <td>{esc(r.subject)}</td>
                <td>{esc(r.assessment)}</td>
                <td>{r.marks:g} / {r.maximum:g}</td>
                <td>{r.result_date.strftime('%d %b %Y')}</td>
                <td>
                    <details>
                        <summary>Edit marks</summary>
                        <form method="post"
                              action="/teacher/learner/{student.id}/results/{r.id}/edit"
                              style="display:grid; gap:8px; min-width:200px;">
                            <input type="hidden" name="csrf_token"
                                   value="{esc(csrf(request))}">

                            <label>Subject
                                <input name="subject" maxlength="120"
                                       value="{esc(r.subject)}" required>
                            </label>

                            <label>Assessment
                                <input name="assessment" maxlength="120"
                                       value="{esc(r.assessment)}" required>
                            </label>

                            <label>Date
                                <input type="date" name="result_date"
                                       value="{r.result_date.isoformat()}"
                                       required>
                            </label>

                            <label>Marks obtained
                                <input type="number" name="marks"
                                       min="0" max="{r.maximum:g}"
                                       step="0.01" value="{r.marks:g}"
                                       required>
                            </label>

                            <label>Maximum marks
                                <input type="number" name="maximum"
                                       min="0.01" step="0.01"
                                       value="{r.maximum:g}" required>
                            </label>

                            <button type="submit">Save changes</button>
                        </form>
                    </details>
                </td>
            </tr>
            """
            for r in results
        ) or (
            '<tr><td colspan="5">'
            'No assessment results recorded.</td></tr>'
        )

        remarks = db.scalars(
            select(TeacherRemark)
            .where(TeacherRemark.student_id == student.id)
            .order_by(TeacherRemark.created_at.desc())
        ).all()
        
        remark_cards = "".join(
            f"""
            <article style="padding:12px; margin:12px 0;
                            border:1px solid var(--border);
                            border-radius:10px;">
                <strong>{esc(r.category)}</strong>
                <p>{esc(r.comment).replace(chr(10), '<br>')}</p>
                <small>
                    By {esc(r.teacher.name or r.teacher.email)}
                    · {r.created_at.strftime('%d %b %Y, %H:%M')}
                </small>
            </article>
            """
            for r in remarks
        ) or "<p>No teacher remarks recorded yet.</p>"

        body = f"""
        <p><a href="/teacher/learners">← Back to all learners</a></p>

        <h1>{esc(student.name)}</h1>
        <p>Grade: {student.grade}</p>
        <p>Admission number: {esc(student.admission or "—")}</p>

        <section>
            <h2>Assessment results</h2> {graph_svg}
            <div style="overflow-x:auto;">
                <table>
                    <thead>
                        <tr>
                            <th>Subject</th>
                            <th>Assessment</th>
                            <th>Marks</th>
                            <th>Date</th>
                            <th>Edit</th>
                        </tr>
                    </thead>
                    <tbody>{result_rows}</tbody>
                </table>
            </div>
        </section>

        <section id="teacher-remarks">
            <h2>Teacher Remarks</h2>
            <p>Record this learner's progress, effort, and areas for improvement.</p>

            <form method="post"
                  action="/teacher/learner/{student.id}/remarks">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">

                <label>Remark category
                    <select name="category" required>
                        <option value="Academic progress">Academic progress</option>
                        <option value="Behaviour and discipline">Behaviour and discipline</option>
                        <option value="Class participation">Class participation</option>
                        <option value="Effort and improvement">Effort and improvement</option>
                        <option value="General observation">General observation</option>
                    </select>
                </label>

                <label>Teacher's comment
                    <textarea name="comment" rows="4" maxlength="5000"
                              required style="width:100%;"
                              placeholder="Write your observation here..."></textarea>
                </label>

                <button type="submit">Save remark</button>
            </form>

            <hr>
            <h3>Previous remarks</h3>
            {remark_cards}
        </section>
        """

        return page(request, "Learner Profile", body, user)



@app.get("/parent/learner/{student_id}", response_class=HTMLResponse)
def parent_learner_profile(request: Request, student_id: int):
    with DB() as db:
        user = require_role(request, db, "parent")

        # Security: parents can only view their own linked learners.
        student = next(
            (child for child in user.children if child.id == student_id),
            None,
        )

        if student is None:
            raise HTTPException(
                status_code=404,
                detail="Learner not found.",
            )

        # Sort results chronologically for the graph.
        chart_results = sorted(
            student.results,
            key=lambda r: (r.result_date, r.id),
        )

        graph_points = []
        graph_circles = []
        graph_dates = []

        for i, result in enumerate(chart_results):
            score = (
                result.marks / result.maximum * 100
                if result.maximum
                else 0
            )

            x = 70 + i * (
                650 / max(1, len(chart_results) - 1)
            )
            y = 225 - score * 1.9

            graph_points.append(f"{x:.1f},{y:.1f}")

            tooltip = (
                f"{esc(result.subject)} — "
                f"{esc(result.assessment)}: {score:.1f}% "
                f"({result.result_date.strftime('%d %b %Y')})"
            )

            graph_circles.append(
                f'<circle cx="{x:.1f}" cy="{y:.1f}" r="5" '
                f'fill="currentColor">'
                f'<title>{tooltip}</title></circle>'
            )

            graph_dates.append(
                f'<text x="{x:.1f}" y="250" font-size="10" '
                f'text-anchor="end" '
                f'transform="rotate(-35 {x:.1f} 250)">'
                f'{result.result_date.strftime("%d %b %y")}</text>'
            )

        graph_grid = "".join(
            f'<line x1="70" y1="{225 - pct * 1.9}" '
            f'x2="730" y2="{225 - pct * 1.9}" '
            f'stroke="var(--border)" stroke-width="1"/>'
            f'<text x="58" y="{230 - pct * 1.9}" '
            f'text-anchor="end" font-size="11">{pct}%</text>'
            for pct in (0, 25, 50, 75, 100)
        )

        if chart_results:
            graph_svg = f"""
            <div style="width:100%; overflow-x:auto;">
                <svg viewBox="0 0 760 310"
                     role="img"
                     aria-label="Learner assessment performance over time"
                     style="width:100%; min-width:420px;
                            max-height:360px; color:var(--text);">
                    {graph_grid}
                    <line x1="70" y1="35" x2="70" y2="225"
                          stroke="currentColor"/>
                    <line x1="70" y1="225" x2="730" y2="225"
                          stroke="currentColor"/>
                    <polyline points="{' '.join(graph_points)}"
                              fill="none" stroke="currentColor"
                              stroke-width="2.5"
                              stroke-linejoin="round"
                              stroke-linecap="round"/>
                    {''.join(graph_circles)}
                    {''.join(graph_dates)}
                    <text x="400" y="305" text-anchor="middle"
                          font-size="12">Assessment date</text>
                </svg>
            </div>
            """
        else:
            graph_svg = (
                "<p>No assessment results available to graph yet.</p>"
            )

        # Display results without teacher editing controls.
        results = sorted(
            student.results,
            key=lambda r: (r.result_date, r.subject),
            reverse=True,
        )

        result_rows = "".join(
    f"""
    <tr>
        <td>{esc(r.subject)}</td>
        <td>{esc(r.assessment)}</td>
        <td>{r.result_date.strftime('%d %b %Y')}</td>
        <td>{r.marks:g} / {r.maximum:g}</td>
        <td>
            {f"{r.marks / r.maximum * 100:.1f}%" if r.maximum > 0 else "Invalid maximum"}
        </td>
    </tr>
    """
    for r in results
) or """
            <tr>
                <td colspan="5">No assessment results recorded yet.</td>
            </tr>
        """

        # Retrieve this learner's teacher remarks.
        remarks = db.scalars(
            select(TeacherRemark)
            .where(TeacherRemark.student_id == student.id)
            .order_by(TeacherRemark.created_at.desc())
        ).all()

        remark_cards = "".join(
            f"""
            <article style="padding:12px; margin:12px 0;
                            border:1px solid var(--border);
                            border-radius:10px;">
                <strong>{esc(r.category)}</strong>
                <p>{esc(r.comment).replace(chr(10), '<br>')}</p>
                <small>
                    Teacher: {esc(r.teacher.name or r.teacher.email)}
                    · {r.created_at.strftime('%d %b %Y, %H:%M')}
                </small>
            </article>
            """
            for r in remarks
        ) or "<p>No teacher remarks recorded yet.</p>"

        body = f"""
        <p><a href="/parent">← Back to Parent Portal</a></p>

        <h1>{esc(student.name)}</h1>
        <p>Grade: {student.grade}</p>
        <p>Admission number: {esc(student.admission or "—")}</p>

        <section>
            <h2>Performance graph</h2>
            <p>Assessment percentages over time.</p>
            {graph_svg}
        </section>

        <section>
            <h2>Assessment results</h2>
            <div style="overflow-x:auto;">
                <table>
                    <thead>
                        <tr>
                            <th>Subject</th>
                            <th>Assessment</th>
                            <th>Date</th>
                            <th>Marks</th>
                            <th>Percentage</th>
                        </tr>
                    </thead>
                    <tbody>{result_rows}</tbody>
                </table>
            </div>
        </section>

        <section id="teacher-remarks">
            <h2>Teacher remarks</h2>
            {remark_cards}
        </section>
        """

        return page(request, "Learner Profile", body, user)



@app.post("/teacher/learner/{student_id}/remarks")
def add_teacher_remark(
    request: Request,
    student_id: int,
    category: str = Form(...),
    comment: str = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    allowed_categories = {
        "Academic progress",
        "Behaviour and discipline",
        "Class participation",
        "Effort and improvement",
        "General observation",
    }

    category = category.strip()
    comment = comment.strip()

    if category not in allowed_categories:
        raise HTTPException(
            status_code=400,
            detail="Please select a valid remark category.",
        )

    if not comment or len(comment) > 5000:
        raise HTTPException(
            status_code=400,
            detail="Remark must contain 1–5000 characters.",
        )

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        student = db.scalar(
            select(Student).where(
                Student.id == student_id,
                Student.school_id == teacher.school_id,
            )
        )

        if not student:
            raise HTTPException(
                status_code=404,
                detail="Learner not found.",
            )

        remark = TeacherRemark(
            student_id=student.id,
            teacher_id=teacher.id,
            category=category,
            comment=comment,
        )

        db.add(remark)
        db.commit()

    return go(f"/teacher/learner/{student_id}#teacher-remarks")

# ============================================================
# 6. REGISTER LEARNERS AND PARENTS
# ============================================================

@app.post("/students")
def add_student(
    request: Request,
    name: str = Form(...),
    grade: int = Form(...),
    admission: str = Form(""),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        name = name.strip()
        admission = admission.strip() or None

        if not name:
            raise HTTPException(400, "Learner name is required.")

        if grade not in GRADES:
            raise HTTPException(400, "Grade must be between 1 and 12.")

        if admission and db.scalar(
            select(Student).where(
                Student.admission == admission,
                Student.school_id == teacher.school_id,
            )
        ):
            raise HTTPException(400, "That admission number already exists.")

        db.add(
            Student(
                name=name,
                grade=grade,
                admission=admission,
                school_id=teacher.school_id,
            )
        )
        db.commit()

    return go("/teacher/add-learner")


@app.post("/parents")
def add_parent(
    request: Request,
    email: str = Form(...),
    password: str = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        email = email.strip().lower()

        if len(password) < 12:
            raise HTTPException(
                400, "Password must be at least 12 characters."
            )

        if db.scalar(select(User).where(User.email == email)):
            raise HTTPException(
                400, "An account with that email already exists."
            )

        db.add(
            User(
                email=email,
                password_hash=passwords.hash(password),
                role="parent",
                school_id=teacher.school_id,
            )
        )
        db.commit()

    return go("/teacher/add-learner")


# ============================================================
# 7. RESET PARENT PASSWORD
# ============================================================

@app.post("/parents/reset-password", response_class=HTMLResponse)
def reset_parent_password(
    request: Request,
    parent_id: int = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    with DB() as db:
        teacher = require_role(request, db, "teacher")
        parent = db.get(User, parent_id)

        if not parent or parent.role != "parent":
            body = """
          <section>
            <h1>Parent not found</h1>
            <a class="button" href="/teacher/add-learner">Back</a>
            </section>
            """
            return page(
                request, "Parent not found", body,
                teacher, status_code=404
            )

        temporary_password = secrets.token_urlsafe(12)
        parent.password_hash = passwords.hash(temporary_password)

        try:
            db.commit()
        except Exception:
            db.rollback()
            raise

        # Show the password only after the database save succeeds.
        # It is not stored in readable form in the database.
        body = f"""
        <section>
            <h1>Parent password reset</h1>
            <p>The password has been changed successfully.</p>
            <p><b>Parent email:</b> {esc(parent.email)}</p>
            <p><b>New temporary password:</b></p>
            <p><code id="temporary-password">
                {esc(temporary_password)}
            </code></p>

            <button type="button" id="copy-password">
                Copy password
            </button>

            <p class="note">
                Copy this password now and share it privately
                with the correct parent. This page will not
                display the password again after you leave it.
            </p>

            <p>
                <a class="button" href="/teacher/add-learner">
                    Return to Add Learner
                </a>
            </p>
        </section>

        <script>
        document.getElementById("copy-password").addEventListener(
            "click", async function () {{
                const password = document
                    .getElementById("temporary-password")
                    .textContent.trim();

                try {{
                    await navigator.clipboard.writeText(password);
                    this.textContent = "Copied";
                }} catch (error) {{
                    alert("Please select and copy the password manually.");
                }}
            }}
        );
        </script>
        """

        return page(request, "Parent password reset", body, teacher)


# ============================================================
# 8. LINK PARENTS TO LEARNERS
# ============================================================

@app.post("/link")
def link_parent(
    request: Request,
    parent_id: int = Form(...),
    student_id: int = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        parent = db.get(User, parent_id)
        student = db.get(Student, student_id)

        if not parent or parent.role != "parent" or not student:
            raise HTTPException(
                400, "Select a valid parent and learner."
            )

        if student not in parent.children:
            parent.children.append(student)
            db.commit()

            return go("/teacher/add-learner")



# ============================================================
# 9. RECORD ASSESSMENT MARKS
# ============================================================

@app.post("/results")
def add_result(
    request: Request,
    student_id: int = Form(...),
    subject: str = Form(...),
    assessment: str = Form(...),
    result_date: str = Form(...),
    marks: float = Form(...),
    maximum: float = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        # Only allow the teacher to record results for
        # learners belonging to the teacher's own school.
        student = db.scalar(
            select(Student).where(
                Student.id == student_id,
                Student.school_id == teacher.school_id,
            )
        )

        assessment = assessment.strip()

        if (
            not student
            or subject not in SUBJECTS
            or not assessment
            or len(assessment) > 120
        ):
            raise HTTPException(
                status_code=400,
                detail="Select a valid learner, subject and assessment.",
            )

        if (
            not math.isfinite(marks)
            or not math.isfinite(maximum)
            or marks < 0
            or maximum <= 0
            or marks > maximum
        ):
            raise HTTPException(
                status_code=400,
                detail="Marks must be between zero and maximum marks.",
            )

        try:
            parsed_date = date.fromisoformat(result_date)
        except ValueError:
            raise HTTPException(
                status_code=400,
                detail="Enter a valid date.",
            )

        duplicate = db.scalar(
            select(Result).where(
                Result.student_id == student.id,
                Result.subject == subject,
                Result.assessment == assessment,
                Result.result_date == parsed_date,
            )
        )

        if duplicate:
            raise HTTPException(
                status_code=400,
                detail="This result already exists.",
            )

        db.add(
            Result(
                student_id=student.id,
                subject=subject,
                assessment=assessment,
                result_date=parsed_date,
                marks=marks,
                maximum=maximum,
            )
        )
        db.commit()

        return go("/teacher/add-learner")

# ============================================================
# 10. PARENT DASHBOARD
# ============================================================

@app.get("/parent", response_class=HTMLResponse)
def parent_dashboard(request: Request):
    with DB() as db:
        user = require_role(request, db, "parent")

        body = (
            f"<h1>Parent portal</h1>"
            f"<p>Results for {esc(user.email)}.</p>"
        )

        if not user.children:
            body += """
            <section>
                <p>No learner is linked to your account yet.
                Please contact the school.</p>
            </section>
            """

        for student in user.children:
            results = sorted(
                student.results,
                key=lambda r: (r.result_date, r.subject),
                reverse=True,
            )

            rows = "".join(
                f"<tr><td>{esc(r.subject)}</td>"
                f"<td>{esc(r.assessment)}</td>"
                f"<td>{r.result_date}</td>"
                f"<td>{r.marks:g}/{r.maximum:g}</td>"
                f"<td>{r.marks / r.maximum * 100:.1f}%</td></tr>"
                for r in results
            )

            body += f"""
            <section>
                
                <h2>{esc(student.name)} (Grade {student.grade})</h2>
                <p>
    <a href="/parent/learner/{student.id}">
        View full profile, performance graph and teacher remarks →
    </a>
</p>
                <p>Admission: {esc(student.admission or "—")}</p>
                <div class="table" style="overflow-x:auto; width:100%;">
                    <table>
                        <tr><th>Subject</th><th>Assessment</th>
                            <th>Date</th><th>Marks</th><th>Percentage</th></tr>
                        {rows or '<tr><td colspan="5">No results yet.</td></tr>'}
                    </table>
                </div>
            </section>
            """

        return page(request, "Parent dashboard", body, user)


# ============================================================
# GENERAL APP SETTINGS (theme preference is stored in this browser)
# ============================================================

@app.get("/settings", response_class=HTMLResponse)
def general_settings_page(request: Request):
    with DB() as db:
        user = current_user(request, db)
        if not user:
            return go("/login")

        body = f"""
        <h1>Settings</h1>
        <p>Personalise how EduTrack Kenya looks on this device.</p>
        <section>
            <h2>Appearance</h2>
            <div class="theme-row">
                <div>
                    <strong>Light / dark theme</strong>
                    <p class="theme-status" id="theme-status">Light theme is the default. Your choice is saved in this browser.</p>
                </div>
                <button type="button" id="theme-toggle">Switch to dark theme</button>
            </div>
            <p class="note">This setting changes the appearance only. It does not change your school records or account settings.</p>
        </section>
        {('<section><h2>Excel Import</h2><p>Configure the worksheet names used when importing school data.</p><a class="button" href="/import/settings">Open Excel Import Settings</a></section>' if user.role == 'teacher' else '')}
        <p><a href="{('/teacher' if user.role == 'teacher' else '/parent')}">← Back to dashboard</a></p>
        """
        return page(request, "Settings", body, user)


# ============================================================
# 11. EXCEL WORKSHEET SETTINGS
# ============================================================

def worksheet_settings(db):
    return {
        key: get_setting(db, key, default)
        for key, default in DEFAULT_SHEETS.items()
    }


@app.get("/import/settings", response_class=HTMLResponse)
def import_settings_page(request: Request):
    with DB() as db:
        teacher = require_role(request, db, "teacher")
        names = worksheet_settings(db)

        fields = [
            ("learners_sheet", "Learners worksheet"),
            ("parents_sheet", "Parents worksheet"),
            ("links_sheet", "Parent links worksheet"),
            ("marks_sheet", "Assessment marks worksheet"),
        ]

        inputs = "".join(
            f'<label>{label}<input name="{key}" '
            f'value="{esc(names[key])}" maxlength="80" required></label>'
            for key, label in fields
        )

        body = f"""
        <h1>Excel Import Settings</h1>
        <section>
            <p>Change these names to match the worksheet tabs
            in your Excel workbook.</p>
            <form class="grid" method="post"
                  action="/import/settings/save">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">
                {inputs}
                <button type="submit">Save worksheet names</button>
            </form>
        </section>

        <section>
            <h2>Required column headings</h2>
            <p><b>Learners:</b> admission, name, grade</p>
            <p><b>Parents:</b> email, parent_name (optional)</p>
            <p><b>Parent links:</b> parent_email, admission</p>
            <p><b>Assessment marks:</b> admission, subject,
            assessment, date, marks, maximum</p>
            <p class="note">
                Put each subject mark on a separate row.
                A learner can have seven or more subjects.
            </p>
            <a class="button" href="/import">Open Excel import</a>
        </section>
        """

        return page(request, "Excel Import Settings", body, teacher)


@app.post("/import/settings/save")
def save_import_settings(
    request: Request,
    learners_sheet: str = Form(...),
    parents_sheet: str = Form(...),
    links_sheet: str = Form(...),
    marks_sheet: str = Form(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    values = {
        "learners_sheet": learners_sheet.strip(),
        "parents_sheet": parents_sheet.strip(),
        "links_sheet": links_sheet.strip(),
        "marks_sheet": marks_sheet.strip(),
    }

    if any(not v or len(v) > 80 for v in values.values()):
        raise HTTPException(
            400, "Worksheet names must contain 1–80 characters."
        )

    if len({v.casefold() for v in values.values()}) != len(values):
        raise HTTPException(
            400, "Worksheet names must be different."
        )

    with DB() as db:
        teacher = require_role(request, db, "teacher")

        for key, value in values.items():
            set_setting(db, key, value)

        db.commit()
    
    return go("/import/settings")


# ============================================================
# 12. EXCEL READING AND VALIDATION
# ============================================================

def cell_text(value):
    return str(value).strip() if value is not None else ""


def normalize_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = cell_text(value)

    if not text:
        raise ValueError("Date is blank.")

    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        pass

    for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    raise ValueError(
        f"Invalid date '{text}'. Use YYYY-MM-DD or DD/MM/YYYY."
    )

# HEADER_ALIASES

# Alternative Excel headings for each standard field.
HEADER_ALIASES = {
    "admission": {
        "admission", "admission no", "admission number",
        "adm no", "adm number", "student id", "learner id",
        "registration number", "registration no",
    },
    "name": {
        "name", "student name", "learner name",
        "full name", "student full name", "learner full name",
    },
    "grade": {
        "grade", "class", "year", "school grade",
        "grade level", "class level",
    },
    "email": {
        "email", "email address", "parent email",
        "parent email address", "guardian email",
        "guardian email address",
    },
    "parent_name": {
        "parent name", "guardian name", "full name",
        "parent full name", "guardian full name",
    },
    "parent_email": {
        "parent email", "parent email address",
        "guardian email", "guardian email address",
    },
    "subject": {
        "subject", "learning area", "learning area name",
        "subject name",
    },
    "assessment": {
        "assessment", "assessment name", "exam",
        "exam name", "test", "test name",
    },
    "date": {
        "date", "exam date", "assessment date",
        "test date", "date of assessment",
    },
    "marks": {
        "marks", "mark", "score", "scores",
        "result", "points", "marks obtained",
    },
    "maximum": {
        "maximum", "maximum marks", "max marks",
        "total marks", "out of", "max score",
        "maximum score",
    },
}


def normalize_header(value):
    """Normalize headings for reliable matching."""
    import re

    text = cell_text(value).strip().casefold()
    text = re.sub(r"[_\-]+", " ", text)
    text = re.sub(r"[^\w\s]", "", text)
    return " ".join(text.split())


def read_records(ws, required_headers):
    iterator = ws.iter_rows(values_only=True)

    try:
        first = next(iterator)
    except StopIteration:
        raise ValueError(
            f"Worksheet '{ws.title}' is empty."
        )

    # Map every supplied heading to a standard field name.
    aliases = {
        normalize_header(alias): standard
        for standard, alternatives in HEADER_ALIASES.items()
        for alias in (alternatives | {standard})
    }

    headers = []
    for cell in first:
        normalized = normalize_header(cell)
        headers.append(aliases.get(normalized, normalized))

    nonempty_headers = [h for h in headers if h]

    if len(nonempty_headers) != len(set(nonempty_headers)):
        raise ValueError(
            f"Worksheet '{ws.title}' has duplicate or "
            "conflicting column headings. Check whether two "
            "columns refer to the same field."
        )

    missing = [
        field for field in required_headers
        if field not in headers
    ]

    if missing:
        raise ValueError(
            f"Worksheet '{ws.title}' is missing required fields: "
            f"{', '.join(missing)}. Accepted alternative headings "
            "are supported; check the first row of your worksheet."
        )

    output = []

    for row_number, row in enumerate(iterator, start=2):
        if not any(cell_text(v) for v in row):
            continue

        item = {
            headers[i]: row[i] if i < len(row) else None
            for i in range(len(headers))
            if headers[i]
        }

        item["_row"] = row_number
        output.append(item)

        if len(output) > MAX_IMPORT_ROWS:
            raise ValueError(
                f"Worksheet '{ws.title}' exceeds "
                f"{MAX_IMPORT_ROWS} data rows."
            )

    return output


def read_wide_marks(ws):
    """Convert one-row-per-learner marks into result records."""
    iterator = ws.iter_rows(values_only=True)

    try:
        first = next(iterator)
    except StopIteration:
        raise ValueError(
            f"Worksheet '{ws.title}' is empty."
        )

    aliases = {
        normalize_header(alias): standard
        for standard, alternatives in HEADER_ALIASES.items()
        for alias in alternatives | {standard}
    }

    headers = [
        aliases.get(normalize_header(value), normalize_header(value))
        for value in first
    ]

    nonempty = [h for h in headers if h]

    if len(nonempty) != len(set(nonempty)):
        raise ValueError(
            f"Worksheet '{ws.title}' has duplicate or "
            "conflicting column headings."
        )

    required = {"admission", "assessment", "date"}
    missing = required - set(headers)

    if missing:
        raise ValueError(
            f"Wide-format marks need these columns: "
            f"{', '.join(sorted(missing))}."
        )

    metadata = {
        "admission", "assessment", "date",
        "maximum", "marks", "subject",
    }

    subject_columns = [
        (i, header)
        for i, header in enumerate(headers)
        if header and header not in metadata
    ]

    if not subject_columns:
        raise ValueError(
            f"Worksheet '{ws.title}' has no subject columns."
        )

    output = []

    for row_number, row in enumerate(iterator, start=2):
        if not any(cell_text(value) for value in row):
            continue

        values = {
            headers[i]: row[i] if i < len(row) else None
            for i in range(len(headers))
            if headers[i]
        }

        admission = values.get("admission")
        assessment = values.get("assessment")
        result_date = values.get("date")

        if not cell_text(admission):
            raise ValueError(
                f"Row {row_number}: admission number is missing."
            )

        if not cell_text(assessment) or result_date is None:
            raise ValueError(
                f"Row {row_number}: assessment and date are required."
            )

        for column_index, subject in subject_columns:
            raw_mark = (
                row[column_index]
                if column_index < len(row)
                else None
            )

            # Blank subject cells mean no result was supplied.
            if raw_mark is None or not cell_text(raw_mark):
                continue

            output.append({
                "admission": admission,
                "subject": subject.title(),
                "assessment": assessment,
                "date": result_date,
                "marks": raw_mark,
                "maximum": (
    values["maximum"]
    if values.get("maximum") is not None
    and cell_text(values.get("maximum"))
    else 100
),
                "_row": row_number,
            })

            if len(output) > MAX_IMPORT_ROWS:
                raise ValueError(
                    f"Converted marks exceed {MAX_IMPORT_ROWS} rows."
                )

    return output

def read_import_workbook_mode(contents, names, mode):
    try:
        workbook = load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            f"Could not open this Excel workbook: {exc}"
        )

    try:
        learners = []
        parents = []
        links = []
        marks = []

        if mode == "marks_only":
            required_sheets = [names["marks_sheet"]]
        elif mode == "people_only":
            required_sheets = [
                names["learners_sheet"],
                names["parents_sheet"],
                names["links_sheet"],
            ]
        elif mode == "full":
            required_sheets = list(names.values())
        else:
            raise ValueError("Please select a valid import type.")

        missing = [
            sheet for sheet in required_sheets
            if sheet not in workbook.sheetnames
        ]

        if missing:
            raise ValueError(
                "Missing worksheet(s): "
                + ", ".join(missing)
                + ". Check Excel Import Settings."
            )

        if mode in ("people_only", "full"):
            learners = read_records(
                workbook[names["learners_sheet"]],
                ["admission", "name", "grade"],
            )
            parents = read_records(
                workbook[names["parents_sheet"]],
                ["email"],
            )
            links = read_records(
                workbook[names["links_sheet"]],
                ["parent_email", "admission"],
            )

        if mode in ("marks_only", "full"):
            marks_ws = workbook[names["marks_sheet"]]
            first_row = next(
                marks_ws.iter_rows(
                    min_row=1, max_row=1, values_only=True
                ),
                (),
            )
            aliases = {
                normalize_header(alias): standard
                for standard, alternatives in HEADER_ALIASES.items()
                for alias in alternatives | {standard}
            }
            normalized_headers = [
                aliases.get(
                    normalize_header(value),
                    normalize_header(value),
                )
                for value in first_row
            ]
            long_format_headers = {
                "admission", "subject", "assessment",
                "date", "marks", "maximum",
            }

            if long_format_headers.issubset(set(normalized_headers)):
                marks = read_records(
                    marks_ws,
                    [
                        "admission", "subject", "assessment",
                        "date", "marks", "maximum",
                    ],
                )
            else:
                marks = read_wide_marks(marks_ws)

        return learners, parents, links, marks

    finally:
        workbook.close()


def read_import_workbook(contents, names):
    try:
        workbook = load_workbook(
            BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            f"Could not open this Excel workbook: {exc}"
        ) from exc

    try:
        missing = [
            name for name in names.values()
            if name not in workbook.sheetnames
        ]

        if missing:
            raise ValueError(
                "Missing worksheet(s): "
                + ", ".join(missing)
                + ". Check Excel Import Settings."
            )

        learners = read_records(
            workbook[names["learners_sheet"]],
            ["admission", "name", "grade"],
        )

        parents = read_records(
            workbook[names["parents_sheet"]],
            ["email"],
        )

        links = read_records(
            workbook[names["links_sheet"]],
            ["parent_email", "admission"],
        )

        marks_ws = workbook[names["marks_sheet"]]

        # Inspect the marks headings to select the right format.
        first_row = next(
            marks_ws.iter_rows(
                min_row=1, max_row=1, values_only=True
            ),
            (),
        )

        aliases = {
            normalize_header(alias): standard
            for standard, alternatives in HEADER_ALIASES.items()
            for alias in alternatives | {standard}
        }

        normalized_headers = [
            aliases.get(
                normalize_header(value),
                normalize_header(value),
            )
            for value in first_row
        ]

        long_required = {
            "admission", "subject", "assessment",
            "date", "marks", "maximum",
        }

        if long_required.issubset(set(normalized_headers)):
            marks = read_records(
                marks_ws,
                [
                    "admission", "subject", "assessment",
                    "date", "marks", "maximum",
                ],
            )
        else:
            marks = read_wide_marks(marks_ws)

        return learners, parents, links, marks

    finally:
        workbook.close()


def validate_workbook(db, learners, parents, links, marks, school_id):
    errors = []
    learner_admissions = set()
    parent_emails = set()
    seen_marks = set()

    # Validate learner rows.
    for row in learners:
        admission = cell_text(row.get("admission"))
        name = cell_text(row.get("name"))
        try:
            grade_value = float(row.get("grade"))

            if not math.isfinite(grade_value) or not grade_value.is_integer():
                raise ValueError

            grade = int(grade_value)

        except (TypeError, ValueError, OverflowError):
            errors.append(
                f"Learners row {row['_row']}: "
                "grade must be a whole number between 1 and 12."
            )
            continue

        if grade not in GRADES:
            errors.append(
                f"Learners row {row['_row']}: "
                "grade must be between 1 and 12."
            )
            continue

        if grade not in GRADES:
            errors.append(
                f"Learners row {row['_row']}: "
                "grade must be between 1 and 12."
            )

        if admission in learner_admissions:
            errors.append(
                f"Learners row {row['_row']}: "
                f"duplicate admission '{admission}'."
            )

        learner_admissions.add(admission)

        if admission and db.scalar(
            select(Student).where(Student.admission == admission)
        ):
            errors.append(
                f"Learners row {row['_row']}: "
                f"admission '{admission}' already exists."
            )

    # Validate parent rows.
    for row in parents:
        email = cell_text(row.get("email")).lower()

        if (
            "@" not in email
            or "." not in email.split("@")[-1]
            or len(email) > 255
        ):
            errors.append(
                f"Parents row {row['_row']}: invalid email '{email}'."
            )

        if email in parent_emails:
            errors.append(
                f"Parents row {row['_row']}: "
                f"duplicate email '{email}'."
            )

        parent_emails.add(email)

        if email and db.scalar(
            select(User).where(User.email == email)
        ):
            errors.append(
                f"Parents row {row['_row']}: "
                f"email '{email}' already exists."
            )

    # Validate parent-to-learner links.
    for row in links:
        email = cell_text(row.get("parent_email")).lower()
        admission = cell_text(row.get("admission"))

        parent_exists = (
            email in parent_emails
            or db.scalar(
                select(User).where(
                    User.email == email,
                    User.role == "parent",
                    User.school_id == school_id
                )
            ) is not None
        )

        student_exists = (
            admission in learner_admissions
            or db.scalar(
                select(Student).where(
                    Student.admission == admission,
                    Student.school_id == school_id
                )
            ) is not None
        )

        if not parent_exists:
            errors.append(
                f"ParentLinks row {row['_row']}: "
                f"parent '{email}' was not found."
            )

        if not student_exists:
            errors.append(
                f"ParentLinks row {row['_row']}: "
                f"admission '{admission}' was not found."
            )

    # Validate assessment marks.
    for row in marks:
        admission = cell_text(row.get("admission"))
        subject = cell_text(row.get("subject"))
        assessment = cell_text(row.get("assessment"))

        student_exists = (
            admission in learner_admissions
            or db.scalar(
                select(Student).where(
                    Student.admission == admission,
                    Student.school_id == school_id
                )
            ) is not None
        )

        if not student_exists:
            errors.append(
                f"AssessmentMarks row {row['_row']}: "
                f"admission '{admission}' was not found."
            )

        if subject not in SUBJECTS:
            errors.append(
                f"AssessmentMarks row {row['_row']}: "
                f"unknown subject '{subject}'."
            )

        if not assessment or len(assessment) > 120:
            errors.append(
                f"AssessmentMarks row {row['_row']}: "
                "assessment is required and max 120 characters."
            )

        try:
            score = float(row.get("marks"))
            maximum = float(row.get("maximum"))

            if not math.isfinite(score) or not math.isfinite(maximum):
                raise ValueError("Marks and maximum must be finite numbers.")

            if maximum <= 0:
                raise ValueError("Maximum mark must be greater than zero.")

            if score < 0:
                raise ValueError("Mark cannot be negative.")

            if score > maximum:
                raise ValueError(
                    f"Mark {score:g} exceeds maximum {maximum:g}."
                )

        except (TypeError, ValueError, OverflowError) as exc:
            errors.append(
                f"AssessmentMarks row {row['_row']}: {exc}"
            )
            continue

        try:
            parsed_date = normalize_date(row.get("date"))
        except ValueError as exc:
            errors.append(
                f"AssessmentMarks row {row['_row']}: {exc}"
            )
            continue

        key = (
            admission,
            subject,
            assessment.casefold(),
            parsed_date.isoformat(),
        )

        if key in seen_marks:
            errors.append(
                f"AssessmentMarks row {row['_row']}: "
                "duplicate mark for learner, subject, assessment and date."
            )

        seen_marks.add(key)

    return errors


# ============================================================
# 13. EXCEL IMPORT PAGE
# ============================================================

@app.get("/import", response_class=HTMLResponse)
def import_page(request: Request):
    with DB() as db:
        user = require_role(request, db, "teacher")
        names = worksheet_settings(db)

        body = f"""
        <h1>Import learners, parents and marks</h1>
        <p>Upload an Excel workbook (.xlsx). It is validated before saving.</p>

        <section>
            <h2>Configured worksheet tabs</h2>
            <ul>
                <li>Learners: {esc(names['learners_sheet'])}</li>
                <li>Parents: {esc(names['parents_sheet'])}</li>
                <li>Parent links: {esc(names['links_sheet'])}</li>
                <li>Assessment marks: {esc(names['marks_sheet'])}</li>
            </ul>
            <a class="button" href="/import/settings">
                Change worksheet names
            </a>
        </section>

        <section>
            <form method="post" action="/import/preview"
                  enctype="multipart/form-data">
                <input type="hidden" name="csrf_token"
                       value="{esc(csrf(request))}">
                <label>Select Excel workbook (.xlsx)
                    <input type="file" name="file"
                           accept=".xlsx" required>
                </label><br>
                <button type="submit">Validate and preview</button>
            </form>
        </section>

        <section>
            <h2>Required headings</h2>
            <p><b>Learners:</b> admission, name, grade</p>
            <p><b>Parents:</b> email, optional parent_name</p>
            <p><b>Parent links:</b> parent_email, admission</p>
            <p><b>Assessment marks:</b> admission, subject,
            assessment, date, marks, maximum</p>
            <p class="note">
                Each mark occupies one row. Admission numbers connect
                learners, parents and marks.
            </p>
        </section>
        """

        return page(request, "Spreadsheet Import", body, user)


@app.post("/import/preview", response_class=HTMLResponse)
async def import_preview(
    request: Request,
    file: UploadFile = File(...),
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    try:
        if (
            not file.filename
            or not file.filename.lower().endswith(".xlsx")
        ):
            raise HTTPException(
                400, "Please upload an .xlsx workbook."
            )

        contents = await file.read(MAX_IMPORT_BYTES + 1)

        if len(contents) > MAX_IMPORT_BYTES:
            raise HTTPException(
                413, "Workbook exceeds the 5 MB upload limit."
            )

        with DB() as db:
            user = require_role(request, db, "teacher")
            names = worksheet_settings(db)

            learners, parents, links, marks = read_import_workbook(
                contents, names
            )

            errors = validate_workbook(
                db, learners, parents, links, marks, user.school_id
            )

            if errors:
                listing = "".join(
                    f"<li>{esc(error)}</li>"
                    for error in errors[:100]
                )

                body = f"""
                <section>
                    <h1>Correct these errors</h1>
                    <p>{len(errors)} issue(s) found.
                    Nothing was saved.</p>
                    <ul>{listing}</ul>
                    <a class="button" href="/import">
                        Upload corrected workbook
                    </a>
                </section>
                """

                return page(
                    request, "Import validation",
                    body, user, status_code=400
                )

            payload = {
                "learners": [
                    {
                        "admission": cell_text(row["admission"]),
                        "name": cell_text(row["name"]),
                        "grade": int(float(row["grade"])),
                    }
                    for row in learners
                ],
                "parents": [
                    {
                        "email": cell_text(row["email"]).lower(),
                        "parent_name": cell_text(
                            row.get("parent_name")
                        ),
                    }
                    for row in parents
                ],
                "links": [
                    {
                        "parent_email": cell_text(
                            row["parent_email"]
                        ).lower(),
                        "admission": cell_text(row["admission"]),
                    }
                    for row in links
                ],
                "marks": [
                    {
                        "admission": cell_text(row["admission"]),
                        "subject": cell_text(row["subject"]),
                        "assessment": cell_text(row["assessment"]),
                        "date": normalize_date(
                            row["date"]
                        ).isoformat(),
                        "marks": float(row["marks"]),
                        "maximum": float(row["maximum"]),
                    }
                    for row in marks
                ],
            }

            # Remove this teacher's previous pending preview.
            old_token = request.session.get("import_batch_token")

            if old_token:
                old_batch = db.get(ImportBatch, old_token)

                if old_batch and old_batch.teacher_id == user.id:
                    db.delete(old_batch)

            token = secrets.token_urlsafe(32)

            db.add(
                ImportBatch(
                    token=token,
                    teacher_id=user.id,
                    payload=json.dumps(payload),
                )
            )
            db.commit()

            request.session["import_batch_token"] = token

            body = f"""
            <h1>Review import before saving</h1>
            <p>Nothing has been saved yet.</p>

            <section>
                <ul>
                    <li>Learners: <b>{len(learners)}</b></li>
                    <li>Parent accounts: <b>{len(parents)}</b></li>
                    <li>Parent links: <b>{len(links)}</b></li>
                    <li>Assessment marks: <b>{len(marks)}</b></li>
                </ul>

                <p class="note">
                    New parent accounts receive temporary passwords
                    displayed after a successful import.
                </p>

                <form method="post" action="/import/confirm">
                    <input type="hidden" name="csrf_token"
                           value="{esc(csrf(request))}">
                    <button type="submit">Confirm and import</button>
                </form>

                <form method="post" action="/import/cancel">
                    <input type="hidden" name="csrf_token"
                           value="{esc(csrf(request))}">
                    <button type="submit">Cancel import</button>
                </form>
            </section>
            """

            return page(request, "Review import", body, user)

    except HTTPException:
        raise

    except Exception as exc:
        with DB() as db:
            user = current_user(request, db)

        body = f"""
        <section>
            <h1>Workbook could not be read</h1>
            <p class="error">{esc(exc)}</p>
            <a class="button" href="/import">Back</a>
        </section>
        """

        return page(
            request, "Import error", body, user, status_code=400
        )

    finally:
        await file.close()


# ============================================================
# 14. CANCEL IMPORT
# ============================================================

@app.post("/import/cancel")
def cancel_import(
    request: Request,
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    token = request.session.get("import_batch_token")

    with DB() as db:
        user = require_role(request, db, "teacher")

        if token:
            batch = db.get(ImportBatch, token)

            if batch and batch.teacher_id == user.id:
                db.delete(batch)
                db.commit()

            request.session.pop("import_batch_token", None)

    return go("/import")


# ============================================================
# 15. CONFIRM AND SAVE EXCEL IMPORT
# ============================================================

@app.post("/import/confirm", response_class=HTMLResponse)
def import_confirm(
    request: Request,
    csrf_token: str = Form(...),
):
    check_csrf(request, csrf_token)

    # Keep the token if the import fails, so the user can retry.
    token = request.session.get("import_batch_token")

    with DB() as db:
        user = require_role(request, db, "teacher")

        batch = db.get(ImportBatch, token) if token else None

        if not batch or batch.teacher_id != user.id:
            body = """
            <section>
                <h1>Import preview expired</h1>
                <p>Please upload the workbook again.</p>
                <a class="button" href="/import">Back to import</a>
            </section>
            """

            return page(
                request, "Import expired",
                body, user, status_code=400
            )

        data = json.loads(batch.payload)
        temporary_passwords = []

        try:
            # Recheck duplicates before writing.
            for item in data["learners"]:
                existing = db.scalar(
                    select(Student).where(
                        Student.admission == item["admission"]
                    )
                )

                if existing:
                    raise ValueError(
                        f"Admission {item['admission']} already exists."
                    )

            for item in data["parents"]:
                existing = db.scalar(
                    select(User).where(
                        User.email == item["email"],
                        User.school_id == user.school_id,
                        User.role == "parent"
                    )
                )

                if existing:
                    raise ValueError(
                        f"Parent email {item['email']} already exists."
                    )

            learners_by_admission = {}

            for item in data["learners"]:
                student = Student(
                    name=item["name"],
                    grade=item["grade"],
                    admission=item["admission"],
                    school_id=user.school_id,
                )

                db.add(student)
                db.flush()

                learners_by_admission[item["admission"]] = student

            parents_by_email = {}

            for item in data["parents"]:
                temporary_password = secrets.token_urlsafe(12)

                parent = User(
                    email=item["email"],
                    password_hash=passwords.hash(
                        temporary_password
                    ),
                    role="parent",
                    school_id=user.school_id,
                )

                db.add(parent)
                db.flush()

                parents_by_email[item["email"]] = parent

                temporary_passwords.append(
                    {
                        "email": item["email"],
                        "password": temporary_password,
                    }
                )

            # Link parents to the correct learners.
            for item in data["links"]:
                parent = parents_by_email.get(
                    item["parent_email"]
                )

                if not parent:
                    parent = db.scalar(
                        select(User).where(
                            User.email == item["parent_email"],
                            User.role == "parent",
                            User.school_id == user.school_id
                        )
                    )

                student = learners_by_admission.get(
                    item["admission"]
                )

                if not student:
                    student = db.scalar(
                        select(Student).where(
                            Student.admission == item["admission"],
                            Student.school_id == user.school_id
                        )
                    )

                if not parent or not student:
                    raise ValueError(
                        "A parent or learner could not be found."
                    )

                if student not in parent.children:
                    parent.children.append(student)


            # Save marks. Multiple subjects per learner are supported.
            for item in data["marks"]:
                # Validate the score before saving it.
                try:
                    score = float(item["marks"])
                    maximum = float(item["maximum"])

                    if not (
                        math.isfinite(score)
                        and math.isfinite(maximum)
                        and maximum > 0
                        and 0 <= score <= maximum
                    ):
                        raise ValueError

                except (TypeError, ValueError, OverflowError):
                    raise ValueError(
                        f"Invalid marks for admission "
                        f"{item.get('admission', 'unknown')}: "
                        "marks must be between zero and a "
                        "positive maximum."
                    )

                student = learners_by_admission.get(
                    item["admission"]
                )

                if not student:
                    student = db.scalar(
                        select(Student).where(
                            Student.admission == item["admission"],
                            Student.school_id == user.school_id
                        )
                    )

                if not student:
                    raise ValueError(
                        f"Learner {item['admission']} was not found."
                    )

                assessment_date = date.fromisoformat(item["date"])

                duplicate = db.scalar(
                    select(Result)
                    .join(Student, Result.student_id == Student.id)
                    .where(
                        Result.student_id == student.id,
                        Student.school_id == user.school_id,
                        Result.subject == item["subject"],
                        Result.assessment == item["assessment"],
                        Result.result_date == assessment_date,
                    )
                )


                if duplicate:
                    raise ValueError(
                        f"Duplicate result: {student.name}, "
                        f"{item['subject']}, {item['assessment']} "
                        f"on {item['date']}."
                    )

                db.add(
                    Result(
                        student_id=student.id,
                        subject=item["subject"],
                        assessment=item["assessment"],
                        result_date=assessment_date,
                        marks=score,
                        maximum=maximum,
                    )
                )

            # Delete the preview in the same transaction.
            db.delete(batch)
            db.commit()

            # Clear the token only after a successful commit.
            request.session.pop("import_batch_token", None)

        except Exception as exc:
            db.rollback()

            body = f"""
            <section>
                <h1>Import was not completed</h1>
                <p class="error">{esc(exc)}</p>
                <p>No changes from this import were committed.</p>
                <a class="button" href="/import">
                    Return to import
                </a>
            </section>
            """

            return page(
                request, "Import failed", body, user,
                status_code=400
            )

        # Show generated passwords once after successful commit.
        password_rows = "".join(
            f"<tr><td>{esc(item['email'])}</td>"
            f"<td><code>{esc(item['password'])}</code></td></tr>"
            for item in temporary_passwords
        )

        if not password_rows:
            password_rows = (
                '<tr><td colspan="2">'
                'No new parent accounts were created.</td></tr>'
            )

        body = f"""
        <section>
            <h1>Import completed</h1>
            <p>The records were saved successfully.</p>
            <p><a class="button" href="/teacher">
                Return to teacher dashboard
            </a></p>
        </section>

        <section>
            <h2>New parent temporary passwords</h2>
            <p class="note">
                Copy these passwords now and share each one privately
                with its matching parent. They are not stored in
                readable form and cannot be retrieved later.
            </p>
            <div class="table">
                <table>
                    <tr><th>Parent email</th>
                        <th>Temporary password</th></tr>
                    {password_rows}
                </table>
            </div>
        </section>
        """

        return page(request, "Import completed", body, user)