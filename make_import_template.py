
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()

# Sheet 1: Learners
ws = wb.active
ws.title = "Learners"
ws.append(["admission", "name", "grade"])
ws.append(["JSS001", "Sample Learner A", 7])
ws.append(["JSS002", "Sample Learner B", 8])
ws.append(["JSS003", "Sample Learner C", 9])

# Sheet 2: Parents
ws = wb.create_sheet("Parents")
ws.append(["email", "parent_name"])
ws.append(["parent1@example.com", "Sample Parent A"])
ws.append(["parent2@example.com", "Sample Parent B"])

# Sheet 3: ParentLinks
ws = wb.create_sheet("ParentLinks")
ws.append(["parent_email", "admission"])
ws.append(["parent1@example.com", "JSS001"])
ws.append(["parent1@example.com", "JSS002"])
ws.append(["parent2@example.com", "JSS003"])

# Format the headings and columns
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="245A81"
        )
        cell.alignment = Alignment(horizontal="center")

    for column in sheet.columns:
        width = max(
            len(str(cell.value or ""))
            for cell in column
        ) + 4
        sheet.column_dimensions[
            get_column_letter(column[0].column)
        ].width = min(width, 35)

wb.save("jss_import.xlsx")
print("Created jss_import.xlsx successfully.")
